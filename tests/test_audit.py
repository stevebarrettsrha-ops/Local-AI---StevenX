"""Regression tests for the defects a full audit found: data safety, the
local-only guard, memory, recall, the prompt budget, install picks, and
the workspace."""

from __future__ import annotations

import json
import os
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest import mock

import engine
import fit
import server


class Isolated(unittest.TestCase):
    """Every data file in a temporary folder."""

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.old_cfg = dict(server.cfg)
        self.patches = [
            mock.patch.object(server, "CHATS_PATH", self.dir / "chats.json"),
            mock.patch.object(server, "MEMORY_PATH", self.dir / "memory.md"),
            mock.patch.object(server, "CONFIG_PATH", self.dir / "config.json"),
            mock.patch.object(server, "EMB_PATH", self.dir / "emb.json"),
            mock.patch.object(server, "_emb_index", None),
        ]
        for p in self.patches:
            p.start()
        self.client = server.app.test_client()

    def tearDown(self) -> None:
        for p in reversed(self.patches):
            p.stop()
        server.cfg.clear()
        server.cfg.update(self.old_cfg)
        self.tmp.cleanup()


class ChatFileTests(Isolated):
    def test_a_damaged_history_is_kept_aside_not_overwritten(self) -> None:
        server.CHATS_PATH.write_text('[{"id": "a", "messages": []}, ',
                                     encoding="utf-8")
        server.put_chat({"id": "new", "messages": []})
        kept = list(self.dir.glob("chats.json.unreadable-*"))
        self.assertEqual(len(kept), 1)
        self.assertIn('"id": "a"', kept[0].read_text(encoding="utf-8"))
        self.assertEqual([c["id"] for c in server.read_chats()], ["new"])

    def test_a_locked_file_fails_the_write_instead_of_wiping(self) -> None:
        server.write_chats([{"id": "a", "messages": []}])
        with mock.patch.object(Path, "read_text",
                               side_effect=PermissionError("locked")), \
                mock.patch.object(server.time, "sleep"):
            with self.assertRaises(OSError):
                server.put_chat({"id": "b", "messages": []})
        self.assertEqual([c["id"] for c in server.read_chats()], ["a"])

    def test_deleting_a_chat_forgets_its_recall_entries(self) -> None:
        server.write_chats([{"id": "a", "messages": []},
                            {"id": "b", "messages": []}])
        server.emb_index().update({"a:0": {"chat": "a", "v": [1]},
                                   "b:0": {"chat": "b", "v": [1]}})
        self.client.delete("/api/chats/a")
        self.assertEqual(list(server.emb_index()), ["b:0"])
        self.assertEqual(list(json.loads(server.EMB_PATH.read_text())),
                         ["b:0"])


class LocalOnlyTests(unittest.TestCase):
    def setUp(self) -> None:
        self.client = server.app.test_client()

    def test_a_rebound_hostname_is_refused(self) -> None:
        r = self.client.get("/api/memory",
                            headers={"Host": "attacker.example:7806"})
        self.assertEqual(r.status_code, 403)

    def test_another_sites_page_cannot_change_anything(self) -> None:
        r = self.client.post("/api/workspace/run", json={"command": "x"},
                             headers={"Origin": "https://evil.example"})
        self.assertEqual(r.status_code, 403)

    def test_the_apps_own_page_is_served(self) -> None:
        for host in ("127.0.0.1:%d" % server.PORT, "localhost:%d" % server.PORT):
            r = self.client.get("/api/chats", headers={"Host": host})
            self.assertEqual(r.status_code, 200)
        with mock.patch.object(server, "load_chats", return_value=[]), \
                mock.patch.object(server, "write_chats"), \
                mock.patch.object(server, "emb_forget"):
            r = self.client.delete(
                "/api/chats/x", headers={
                    "Origin": "http://127.0.0.1:%d" % server.PORT})
        self.assertEqual(r.status_code, 200)


class MemoryTests(Isolated):
    def test_the_newest_lines_are_kept_and_sent(self) -> None:
        text = "".join(f"- note {i:05d}\n" for i in range(3000))
        r = self.client.post("/api/memory", json={"text": text})
        self.assertTrue(r.get_json()["trimmed"])
        kept = server.read_memory()
        self.assertTrue(kept.endswith("- note 02999\n"))
        self.assertLessEqual(len(kept), server.MEMORY_CAP)
        sent = server.context_extra("q", "c")
        self.assertIn("note 02999", sent)
        self.assertNotIn("note 00000", sent)

    def test_a_body_that_is_not_an_object_is_a_400(self) -> None:
        self.assertEqual(self.client.post("/api/memory", json=[1]).status_code,
                         400)

    def test_recall_rides_with_the_question_not_the_system_prompt(self) -> None:
        with mock.patch.object(server, "semantic_recall", return_value=None), \
                mock.patch.object(server, "recall_snippets",
                                  return_value=["From x — asked: y"]):
            self.assertNotIn("From x", server.context_extra("q", "c"))
            self.assertIn("From x", server.recall_extra("q", "c"))


class RecallIndexTests(Isolated):
    def test_recall_can_walk_the_index_while_it_grows(self) -> None:
        idx = server.emb_index()
        idx.update({f"c:{i}": {"chat": "c", "title": "t", "q": "q", "a": "a",
                               "v": [1.0, 0.0]} for i in range(2000)})
        stop = threading.Event()

        def grow():
            # the size changes all the time but stays bounded
            n = 0
            while not stop.is_set():
                with server._emb_lock:
                    idx[f"d:{n % 50}"] = {"chat": "d", "v": [1.0, 0.0]}
                    if n % 2:
                        idx.pop(f"d:{(n - 1) % 50}", None)
                n += 1
        t = threading.Thread(target=grow)
        # switch threads often, so a walk over the live dict would collide
        interval = sys.getswitchinterval()
        sys.setswitchinterval(1e-6)
        t.start()
        try:
            with mock.patch.object(server, "ensure_embedder",
                                   return_value=True), \
                    mock.patch.object(server, "index_exchanges"), \
                    mock.patch.object(server.embedder, "embed",
                                      return_value=[[1.0, 0.0]]):
                for _ in range(20):
                    self.assertIsInstance(
                        server.semantic_recall("q", "zzz"), list)
        finally:
            stop.set()
            t.join()
            sys.setswitchinterval(interval)

    def test_one_refused_input_does_not_stop_the_rest(self) -> None:
        server.write_chats([{"id": "c", "title": "t", "messages": [
            {"role": "user", "content": "long " * 10},
            {"role": "assistant", "content": "answer one"},
            {"role": "user", "content": "short"},
            {"role": "assistant", "content": "answer two"}]}])

        def embed(texts):
            if len(texts) > 1 or "long" in texts[0]:
                raise RuntimeError("input is too large")
            return [[0.5, 0.5]]
        with mock.patch.object(server, "ensure_embedder", return_value=True), \
                mock.patch.object(server.embedder, "embed", side_effect=embed):
            server.index_exchanges()
        self.assertEqual(list(server.emb_index()), ["c:2"])
        self.assertTrue(server.EMB_PATH.exists())


class BudgetTests(unittest.TestCase):
    def plan(self, history, keep=1, n_ctx=1000):
        with mock.patch.object(server.server, "n_ctx", return_value=n_ctx), \
                mock.patch.object(server.server, "count_tokens",
                                  side_effect=lambda t: len(t) // 4):
            return server.budget("", history, 0, keep)

    def test_a_continuation_keeps_the_reply_it_continues(self) -> None:
        history = [{"role": "user", "content": "q" * 40},
                   {"role": "assistant", "content": "a" * 3000},
                   {"role": "user", "content": server.CONTINUE_NUDGE}]
        plan = self.plan(history, keep=3)
        roles = [m["role"] for m in plan["history"]]
        self.assertEqual(roles, ["user", "assistant", "user"])
        reply = plan["history"][1]["content"]
        self.assertTrue(reply.startswith("…") and reply.endswith("a" * 200))
        self.assertGreaterEqual(plan["free"], server.REPLY_FLOOR - 64)

    def test_what_is_sent_opens_with_the_person(self) -> None:
        history = [{"role": "user", "content": "x" * 1600},
                   {"role": "assistant", "content": "y" * 400},
                   {"role": "user", "content": "z" * 40}]
        plan = self.plan(history)
        self.assertEqual(plan["history"][0]["role"], "user")


class StreamFailureTests(Isolated):
    def test_a_reply_cut_short_by_an_error_says_so(self) -> None:
        def fake(history, params, **_):
            yield {"delta": "half an ans"}
            raise RuntimeError("connection reset")
        with mock.patch.object(server.server, "ready", return_value=True), \
                mock.patch.object(server.server, "n_ctx", return_value=8192), \
                mock.patch.object(server.server, "count_tokens",
                                  side_effect=lambda t: len(t) // 4), \
                mock.patch.object(server, "context_extra", return_value=""), \
                mock.patch.object(server, "recall_extra", return_value=""), \
                mock.patch.object(server, "index_exchanges"), \
                mock.patch.object(server.server, "chat_stream",
                                  side_effect=fake):
            raw = self.client.post("/api/chat", json={"message": "hi"}) \
                .get_data(as_text=True)
        done = [json.loads(c[5:]) for c in raw.split("\n\n")
                if c.startswith("data:") and '"done"' in c][0]
        self.assertEqual(done["stop"], "error")
        self.assertTrue(done["unfinished"])
        reply = server.read_chats()[0]["messages"][-1]
        self.assertEqual((reply["stop"], reply["unfinished"]), ("error", True))
        self.assertIn("connection reset", reply["error"])


class PickTests(unittest.TestCase):
    def test_best_here_is_the_best_that_fits(self) -> None:
        files = [{"name": f"m-{q}.gguf", "path": f"m-{q}.gguf", "quant": q,
                  "size": s * 1024**3, "split": False, "companion": False}
                 for q, s in (("Q2_K", 3), ("Q4_K_M", 5), ("Q6_K", 6.5),
                              ("Q8_0", 8.5), ("F16", 16))]
        verdict = {3: "fits", 5: "fits", 6.5: "fits", 8.5: "tight",
                   16: "too_big"}
        with mock.patch.object(engine, "hf_files", return_value=files), \
                mock.patch.object(fit, "model_config",
                                  return_value=dict(server.GUESSED_SHAPE)), \
                mock.patch.object(server, "assess_at", side_effect=lambda size,
                                  *a, **k: {"verdict": verdict[
                                      size / 1024**3]}), \
                mock.patch.object(fit, "verdict_text", return_value=""):
            d = server.app.test_client().get("/api/fit?repo=x/y").get_json()
        self.assertEqual(d["recommended"], "m-Q6_K.gguf")

    def test_mmproj_on_disk_is_a_companion_not_a_model(self) -> None:
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(engine, "MODELS_DIR", Path(tmp)):
            Path(tmp, "mmproj-model-f16.gguf").write_bytes(b"x")
            (m,) = engine.local_models()
            self.assertTrue(m["companion"])
            d = server.app.test_client().get("/api/fit/local").get_json()
        self.assertEqual(d["models"][0]["fit"], {})


class BandwidthTests(unittest.TestCase):
    def test_names_match_as_whole_words_and_laptops_get_no_number(self):
        self.assertEqual(fit.bandwidth_of("NVIDIA GeForce RTX 4090"), 1008)
        self.assertEqual(fit.bandwidth_of("NVIDIA RTX A1000"), 0)
        self.assertEqual(fit.bandwidth_of("NVIDIA T400 4GB"), 0)
        self.assertEqual(fit.bandwidth_of("NVIDIA L40S"), 0)
        self.assertEqual(fit.bandwidth_of("NVIDIA GeForce RTX 4090 Laptop GPU"),
                         0)
        self.assertEqual(fit.bandwidth_of("NVIDIA GeForce RTX 4070 SUPER"), 504)


class InstallPickTests(unittest.TestCase):
    NAMES = ["cudart-llama-bin-win-cuda-12.4-x64.zip",
             "llama-b1-bin-macos-arm64.tar.gz",
             "llama-b1-bin-ubuntu-vulkan-x64.tar.gz",
             "llama-b1-bin-ubuntu-x64.tar.gz",
             "llama-b1-bin-win-cpu-arm64.zip", "llama-b1-bin-win-cpu-x64.zip",
             "llama-b1-bin-win-cuda-12.4-x64.zip",
             "llama-b1-bin-win-vulkan-x64.zip", "llama-b1-xcframework.zip"]

    def pick(self, system, machine, vendor=""):
        assets = [{"name": n} for n in self.NAMES]
        with mock.patch.object(engine.platform, "system", return_value=system), \
                mock.patch.object(engine.platform, "machine",
                                  return_value=machine):
            return engine._pick_asset(
                assets, engine.asset_choices({"vendor": vendor}))[0]["name"]

    def test_each_machine_gets_its_own_build(self) -> None:
        self.assertEqual(self.pick("Windows", "AMD64"),
                         "llama-b1-bin-win-cpu-x64.zip")
        self.assertEqual(self.pick("Windows", "AMD64", "nvidia"),
                         "llama-b1-bin-win-cuda-12.4-x64.zip")
        self.assertEqual(self.pick("Windows", "ARM64"),
                         "llama-b1-bin-win-cpu-arm64.zip")
        self.assertEqual(self.pick("Linux", "x86_64"),
                         "llama-b1-bin-ubuntu-x64.tar.gz")
        self.assertEqual(self.pick("Darwin", "arm64"),
                         "llama-b1-bin-macos-arm64.tar.gz")

    def test_a_tarball_unpacks_inside_its_folder_only(self) -> None:
        import io
        import tarfile
        with tempfile.TemporaryDirectory() as tmp:
            bad = Path(tmp, "bad.tar.gz")
            with tarfile.open(bad, "w:gz") as t:
                info = tarfile.TarInfo("../escape.txt")
                info.size = 2
                t.addfile(info, io.BytesIO(b"hi"))
            dest = Path(tmp, "out")
            dest.mkdir()
            with self.assertRaises(Exception):
                engine._unpack(bad, dest)
            self.assertFalse(Path(tmp, "escape.txt").exists())

    def test_a_second_click_gets_the_running_install_back(self) -> None:
        gate = threading.Event()
        first = engine.spawn("download", "voice (test)",
                             lambda t: gate.wait(5))
        second = engine.spawn("download", "voice (test)", lambda t: None)
        gate.set()
        self.assertIs(first, second)


class WorkspaceWriteTests(Isolated):
    def test_line_endings_kept_and_a_bak_link_never_followed(self) -> None:
        ws = self.dir / "ws"
        ws.mkdir()
        (ws / "a.py").write_bytes(b"x = 1\r\n")
        (ws / "b.py").write_bytes(b"x = 1\n")
        outside = self.dir / "outside.txt"
        outside.write_text("safe")
        try:
            os.symlink(outside, ws / "b.py.bak")
        except (OSError, NotImplementedError):
            self.skipTest("no symlinks here")
        with mock.patch.object(server, "workspace_root", return_value=ws):
            for name in ("a.py", "b.py"):
                self.client.post("/api/workspace/file",
                                 json={"path": name, "content": "y = 2\n"})
        self.assertEqual((ws / "a.py").read_bytes(), b"y = 2\r\n")
        self.assertEqual((ws / "b.py").read_bytes(), b"y = 2\n")
        self.assertEqual(outside.read_text(), "safe")


class OfficeExportTests(unittest.TestCase):
    def test_slides_use_title_and_content_and_pictures_fit(self) -> None:
        import io
        from PIL import Image
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(server, "IMAGES_OUT", Path(tmp)):
            for n in ("a.png", "b.png"):
                Image.new("RGB", (512, 512)).save(Path(tmp, n))
            data, _ = server.build_pptx("# One\n- a\n# Two\n- b\n"
                                        "image: a.png\nimage: b.png\n")
        prs = Presentation(io.BytesIO(data))
        for slide in prs.slides:
            self.assertEqual(slide.slide_layout.name, "Title and Content")
            for shape in slide.shapes:
                self.assertGreater(shape.height, 0)
                self.assertLessEqual(shape.top + shape.height,
                                     prs.slide_height)

    def test_sheet_names_leading_zeros_and_chart_axes(self) -> None:
        import io
        import re
        import zipfile
        import openpyxl
        data, _ = server.build_xlsx("# Sheet: Revenue: Q1/Q2\nzip,amount\n"
                                    "02134,10\n90210,12.5\n# Chart: bar S\n")
        ws = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
        self.assertEqual(ws.title, "Revenue- Q1-Q2")
        self.assertEqual((ws["A2"].value, ws["A3"].value), ("02134", 90210))
        z = zipfile.ZipFile(io.BytesIO(data))
        xml = z.read(next(n for n in z.namelist()
                          if "charts/chart" in n)).decode()
        self.assertEqual(re.findall(r'<delete val="(\d)"/>', xml), ["0", "0"])


if __name__ == "__main__":
    unittest.main()
