"""
server.py - Llama Studio backend.

Run:  python server.py        (opens http://127.0.0.1:7806)
"""

from __future__ import annotations

import csv
import html as html_mod
import io
import json
import math
import os
import queue
import re
import signal
import subprocess
import tempfile
import threading
import time
import urllib.parse
import uuid
import webbrowser
from pathlib import Path

import requests
from flask import (Flask, Response, jsonify, request, send_file,
                   send_from_directory)

import engine
import fit

APP_DIR = Path(__file__).resolve().parent
DATA_DIR = APP_DIR / "data"
CHATS_PATH = DATA_DIR / "chats.json"
MEMORY_PATH = DATA_DIR / "memory.md"
CONFIG_PATH = DATA_DIR / "config.json"
WEB_DIR = APP_DIR / "web"
PORT = int(os.environ.get("LLAMA_STUDIO_PORT", "7806"))

app = Flask(__name__, static_folder=None)
# Keep accidental oversized prompts/uploads from exhausting the local process,
# while leaving enough headroom for browser-recorded voice clips.
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024
server = engine.Server()
chats_lock = threading.RLock()
config_lock = threading.RLock()
engine_lock = threading.Lock()
memory_lock = threading.RLock()

DEFAULTS = {
    "hf_token": "", "hf_endpoint": "https://huggingface.co",
    # max_tokens 0 means "as much of the context as is left after the
    # prompt" — a fixed small number is what cuts long answers in half.
    # ctx 0 is Auto: the longest window the fit says costs no GPU layer.
    "ctx": 0, "temperature": 0.7, "top_p": 0.95, "max_tokens": 0,
    # "model": sample the way the model's own generation_config.json says;
    # "manual": the temperature and top-p sliders.
    "sampling": "model",
    # q8 KV is the best default on an 8 GB card: it roughly halves cache use
    # while preserving substantially more room for context and GPU layers.
    "system": "", "kv_bits": 8, "auto_start": True, "plan_for": "",
    "last_model": "", "setup_complete": False, "workspace": "",
    "last_model_config": {}, "last_model_sampling": {}, "last_ctx": 0,
    "run_command": "", "recall": True,
    "reply_limit_migrated": False, "ctx_auto_migrated": False,
}


def load_config() -> dict:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    cfg = dict(DEFAULTS)
    if CONFIG_PATH.exists():
        try:
            cfg.update(json.loads(CONFIG_PATH.read_text(encoding="utf-8")))
        except Exception:
            # The migrations below save straight away: without a copy, a
            # damaged file would take the token and workspace with it.
            keep_aside(CONFIG_PATH)
    if not cfg.get("reply_limit_migrated"):
        # The old 1024-token reply limit was a default nobody chose, and it
        # ended long answers mid-sentence with no explanation. Move that one
        # value to automatic; a number the person picked themselves stands.
        if int(cfg.get("max_tokens") or 0) == 1024:
            cfg["max_tokens"] = 0
        cfg["reply_limit_migrated"] = True
        save_config(cfg)
    if not cfg.get("ctx_auto_migrated"):
        # The same story for the old fixed 8k window: a thinking model can
        # spend all of it deliberating and never write the code. The 8k
        # default moves to Auto; any other size was a choice and stands.
        if int(cfg.get("ctx") or 0) == 8192:
            cfg["ctx"] = 0
        cfg["ctx_auto_migrated"] = True
        save_config(cfg)
    return cfg


def keep_aside(path: Path) -> None:
    """Move an unreadable data file out of the way, kept, before anything
    writes a fresh one over it."""
    try:
        path.replace(path.with_name(
            f"{path.name}.unreadable-{time.strftime('%Y%m%d-%H%M%S')}"))
    except OSError:
        pass


def _atomic_json(path: Path, value) -> None:
    """Durably replace a JSON file instead of exposing a half-written file."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + f".{os.getpid()}.{threading.get_ident()}.tmp")
    try:
        with tmp.open("w", encoding="utf-8") as fh:
            json.dump(value, fh, indent=2)
            fh.flush()
            os.fsync(fh.fileno())
        os.replace(tmp, path)
    finally:
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass


def save_config(value: dict) -> None:
    with config_lock:
        _atomic_json(CONFIG_PATH, value)


cfg = load_config()


# --------------------------------------------------------------------------- #
# chats
# --------------------------------------------------------------------------- #
def load_chats() -> list[dict]:
    """Every chat, for code that writes the file back. A file that cannot
    be read is never taken for an empty history: that turned the next save
    into a single chat, and the whole history was gone. A locked file (an
    antivirus scan on Windows) is retried, then the write fails instead; a
    damaged one is kept aside under another name."""
    with chats_lock:
        for attempt in range(5):
            if not CHATS_PATH.exists():
                return []
            try:
                raw = CHATS_PATH.read_text(encoding="utf-8")
            except OSError:
                if attempt == 4:
                    raise
                time.sleep(0.2)
                continue
            try:
                items = json.loads(raw)
                if isinstance(items, list):
                    return items
            except ValueError:
                pass
            keep_aside(CHATS_PATH)
            return []
        return []


def read_chats() -> list[dict]:
    """Every chat, for code that only reads."""
    try:
        return load_chats()
    except OSError:
        return []


def write_chats(items: list[dict]) -> None:
    with chats_lock:
        _atomic_json(CHATS_PATH, items)


def get_chat(chat_id: str) -> dict | None:
    return next((c for c in read_chats() if c["id"] == chat_id), None)


def put_chat(chat: dict) -> None:
    # Hold one re-entrant lock across the entire read/modify/write operation.
    # Separate locks let simultaneous streams overwrite one another's chats.
    with chats_lock:
        items = load_chats()
        for i, c in enumerate(items):
            if c["id"] == chat["id"]:
                items[i] = chat
                break
        else:
            items.insert(0, chat)
        write_chats(items)


# --------------------------------------------------------------------------- #
# shell
# --------------------------------------------------------------------------- #
@app.get("/")
def index():
    return send_from_directory(WEB_DIR, "index.html")


@app.get("/web/<path:name>")
def web_asset(name: str):
    return send_from_directory(WEB_DIR, name)


LOCAL_NAMES = ("127.0.0.1", "localhost", "[::1]")


def _host_name(host: str) -> str:
    host = (host or "").lower()
    return host.split("]")[0] + "]" if host.startswith("[") \
        else host.split(":")[0]


@app.before_request
def local_only():
    """Only this machine's own page may drive the app. Any web page the
    person visits can send requests to 127.0.0.1, and one that rebinds its
    own hostname to 127.0.0.1 can read the answers too — with a workspace
    endpoint that runs commands, that is code execution. So the Host must
    be a loopback name (a rebinding page arrives under its own), and a
    request that changes anything must not come from another site's page."""
    if _host_name(request.host) not in LOCAL_NAMES:
        return jsonify({"error": "Llama Studio only answers its own "
                                 "page on this machine."}), 403
    origin = request.headers.get("Origin")
    if origin and request.method not in ("GET", "HEAD", "OPTIONS"):
        if origin.rstrip("/").lower() not in {
                f"http://{n}:{PORT}" for n in LOCAL_NAMES}:
            return jsonify({"error": "Refused a request from another "
                                     "site."}), 403
    return None


@app.after_request
def security_headers(response):
    """Safe browser defaults for the loopback UI and its JSON API."""
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "DENY")
    response.headers.setdefault("Referrer-Policy", "no-referrer")
    response.headers.setdefault("Cross-Origin-Resource-Policy", "same-origin")
    response.headers.setdefault("Permissions-Policy", "camera=(), geolocation=()")
    if request.path.startswith("/api/"):
        response.headers.setdefault("Cache-Control", "no-store")
    return response


# --------------------------------------------------------------------------- #
# status / hardware
# --------------------------------------------------------------------------- #
@app.get("/api/status")
def api_status():
    hw = fit.hardware(cfg)
    binary = engine.server_binary()
    return jsonify({
        "hardware": hw,
        "llama_installed": bool(binary),
        "llama_path": str(binary) if binary else "",
        "running": server.alive(), "ready": server.ready(),
        "api_url": server.url + "/v1",
        "model": Path(server.model).name if server.model else "",
        "models": engine.local_models(),
        "models_dir": str(engine.MODELS_DIR),
        "engines": engines_list(),
        "workspace": str(workspace_root() or ""),
        "embed": {"installed": bool(embed_model_file()),
                  "ready": embedder.ready(),
                  "indexed": len(emb_index())},
        "voice": {"installed": bool(engine.whisper_binary()
                                    and voice_model_file()),
                  "ready": whisper.ready()},
        "image": {"installed": bool(engine.sd_binary()
                                    and image_model_file())},
        "config": {k: cfg.get(k) for k in
                   ("ctx", "temperature", "top_p", "max_tokens", "system",
                    "kv_bits", "auto_start", "last_model", "plan_for",
                    "run_command", "sampling")},
        # the window the running server was started with (Auto resolved)
        "ctx_running": server.ctx if server.alive() else 0,
        # the loaded model's own sampling, when it publishes any
        "sampling_profile": loaded_sampling(),
        "catalogue": fit.CATALOGUE,
        "presets": fit.presets(),
        "setup_complete": bool(cfg.get("setup_complete")),
        "tail": server.tail(12),
    })


@app.post("/api/config")
def api_config():
    b = request.get_json(silent=True) or {}
    if not isinstance(b, dict):
        return jsonify({"error": "Configuration must be a JSON object."}), 400
    try:
        clean = {}
        if "ctx" in b:
            clean["ctx"] = int(b["ctx"])
            if clean["ctx"] != 0 and not 512 <= clean["ctx"] <= 262_144:
                raise ValueError("Context must be Auto (0) or between 512 "
                                 "and 262144.")
        if "sampling" in b:
            if b["sampling"] not in ("model", "manual"):
                raise ValueError("Sampling must be 'model' or 'manual'.")
            clean["sampling"] = b["sampling"]
        for key in ("temperature", "top_p"):
            if key in b:
                clean[key] = float(b[key])
                if not math.isfinite(clean[key]):
                    raise ValueError(f"{key} must be a finite number.")
        if "temperature" in clean and not 0 <= clean["temperature"] <= 2:
            raise ValueError("Temperature must be between 0 and 2.")
        if "top_p" in clean and not 0 < clean["top_p"] <= 1:
            raise ValueError("Top-p must be greater than 0 and at most 1.")
        if "max_tokens" in b:
            clean["max_tokens"] = int(b["max_tokens"])
            if not 0 <= clean["max_tokens"] <= 262_144:
                raise ValueError("Reply limit must be between 0 and 262144.")
        if "kv_bits" in b:
            clean["kv_bits"] = int(b["kv_bits"])
            if clean["kv_bits"] not in (8, 16):
                raise ValueError("KV precision must be 8 or 16 bits.")
        for key in ("auto_start", "setup_complete", "recall"):
            if key in b:
                if not isinstance(b[key], bool):
                    raise ValueError(f"{key} must be true or false.")
                clean[key] = b[key]
        for key, limit in (("system", 64_000), ("hf_token", 4_096),
                           ("hf_endpoint", 2_048), ("plan_for", 200)):
            if key in b:
                if not isinstance(b[key], str) or len(b[key]) > limit:
                    raise ValueError(f"{key} is not a valid string.")
                clean[key] = b[key]
        if clean.get("plan_for") and clean["plan_for"] not in fit.GPUS:
            raise ValueError("Unknown planned GPU.")
    except (TypeError, ValueError) as exc:
        return jsonify({"error": str(exc)}), 400
    with config_lock:
        cfg.update(clean)
        save_config(cfg)
    return jsonify({"ok": True})


# --------------------------------------------------------------------------- #
# memory (persistent notes + recall of relevant past conversations)
# --------------------------------------------------------------------------- #
MEMORY_CAP = 16_000        # characters kept in the memory file
MEMORY_INJECT = 4_000      # characters of it sent with each request


def read_memory() -> str:
    with memory_lock:
        try:
            return MEMORY_PATH.read_text(encoding="utf-8") \
                if MEMORY_PATH.exists() else ""
        except OSError:
            return ""


def newest_lines(text: str, limit: int) -> str:
    """The newest whole lines of `text` that fit in `limit` characters.
    Memory grows at the end, so the end is what matters: cutting at a
    character count kept the oldest notes and dropped every new one."""
    if len(text) <= limit:
        return text
    kept, size = [], 0
    for line in reversed(text.splitlines()):
        if size + len(line) + 1 > limit:
            break
        kept.append(line)
        size += len(line) + 1
    return "\n".join(reversed(kept)) + ("\n" if kept else "")


def write_memory(text: str) -> bool:
    """Save the memory file; False when it had to drop its oldest lines
    to stay within MEMORY_CAP."""
    with memory_lock:
        MEMORY_PATH.parent.mkdir(parents=True, exist_ok=True)
        tmp = MEMORY_PATH.with_name(MEMORY_PATH.name + ".tmp")
        kept = newest_lines(text, MEMORY_CAP)
        try:
            tmp.write_text(kept, encoding="utf-8")
            os.replace(tmp, MEMORY_PATH)
        finally:
            tmp.unlink(missing_ok=True)
        return kept == text


def remember_lines(reply: str) -> int:
    """Lines the model marked `remember:` become permanent memory. The
    file is user-editable and capped, so a chatty model cannot flood it."""
    new = [ln.strip()[9:].strip() for ln in reply.splitlines()
           if ln.strip().lower().startswith("remember:")]
    new = [n for n in new if n]
    if not new:
        return 0
    with memory_lock:
        mem = read_memory()
        have = {ln.lstrip("- ").strip().lower() for ln in mem.splitlines()}
        added = [n for n in new if n.lower() not in have]
        if added:
            write_memory((mem.rstrip() + "\n" if mem.strip() else "") +
                         "\n".join("- " + n for n in added) + "\n")
        return len(added)


# ---- semantic recall: a small embedding model beside the chat model ---- #
EMBED_REPO = "nomic-ai/nomic-embed-text-v1.5-GGUF"
EMBED_DIR = "embed"
EMBED_SIM_MIN = 0.5
EMB_PATH = DATA_DIR / "embeddings.json"
embedder = engine.EmbedServer(8082)
# _emb_lock guards the index itself: recall reads it while indexing adds
# to it on another thread, which used to fail the chat with "dictionary
# changed size during iteration". _emb_run lets one indexing pass run at a
# time; a second caller simply skips.
_emb_lock = threading.Lock()
_emb_run = threading.Lock()
_emb_index: dict | None = None


def embed_model_file() -> Path | None:
    d = engine.MODELS_DIR / EMBED_DIR
    if d.is_dir():
        for f in sorted(d.glob("*.gguf")):
            return f
    return None


def emb_index() -> dict:
    global _emb_index
    with _emb_lock:
        if _emb_index is None:
            try:
                _emb_index = json.loads(
                    EMB_PATH.read_text(encoding="utf-8")) \
                    if EMB_PATH.exists() else {}
            except Exception:
                _emb_index = {}
        return _emb_index


def emb_save() -> None:
    with _emb_lock:
        _atomic_json(EMB_PATH, dict(_emb_index or {}))


def emb_entries() -> list[dict]:
    """A snapshot of the index, safe to walk while indexing goes on."""
    emb_index()
    with _emb_lock:
        return list((_emb_index or {}).values())


def emb_forget(chat_id: str) -> None:
    """Drop a deleted conversation from the index, so what it said is not
    recalled into other conversations afterwards."""
    idx = emb_index()
    with _emb_lock:
        gone = [k for k in idx if k.startswith(chat_id + ":")]
        for k in gone:
            del idx[k]
    if gone:
        emb_save()


def ensure_embedder(wait: int = 15) -> bool:
    model = embed_model_file()
    if not model or not engine.server_binary():
        return False
    if embedder.ready():
        return True
    if not embedder.alive():
        try:
            # An embedding model reads its whole input in one micro-batch,
            # and llama-server refuses an input longer than that (512 by
            # default): one long exchange, CJK above all, then failed the
            # batch it sat in on every pass. -ub as large as the window.
            embedder.start(str(model), 0, 2048,
                           ["--embeddings", "-b", "2048", "-ub", "2048"])
        except Exception:
            return False
    return embedder.wait_ready(wait)


def _cos(a: list[float], b: list[float]) -> float:
    num = sum(x * y for x, y in zip(a, b))
    da = sum(x * x for x in a) ** 0.5
    db = sum(y * y for y in b) ** 0.5
    return num / (da * db) if da and db else 0.0


def _exchange_rows():
    for c in read_chats():
        msgs = c.get("messages", [])
        for i, m in enumerate(msgs):
            if m.get("role") != "user":
                continue
            reply = msgs[i + 1]["content"] if i + 1 < len(msgs) \
                and msgs[i + 1].get("role") == "assistant" else ""
            yield c, i, m.get("content", ""), reply


def index_exchanges() -> None:
    """Embed every stored exchange the index does not have yet. The
    nomic prefixes (search_document / search_query) are the model's own
    convention for asymmetric retrieval."""
    if not _emb_run.acquire(blocking=False):
        return
    try:
        if not ensure_embedder():
            return
        idx = emb_index()
        todo = []
        for c, i, um, am in _exchange_rows():
            key = f'{c["id"]}:{i}'
            # An exchange whose reply has not arrived yet is skipped, not
            # indexed half-finished: once a key exists it is never
            # revisited, so indexing early would freeze an empty answer.
            if key in idx or not am.strip():
                continue
            doc = "search_document: " + um[:300].replace("\n", " ") + \
                " " + am[:500].replace("\n", " ")
            todo.append((key, c, um, am, doc))
        added = 0
        for start in range(0, len(todo), 8):
            batch = todo[start:start + 8]
            try:
                vecs = embedder.embed([t[4] for t in batch])
            except Exception:
                # One input the server refuses must not stop the rest:
                # try them one by one and leave out only the failure.
                vecs = []
                for t in batch:
                    try:
                        vecs.append(embedder.embed([t[4]])[0])
                    except Exception:
                        vecs.append(None)
            with _emb_lock:
                for (key, c, um, am, _), v in zip(batch, vecs):
                    if v is None:
                        continue
                    idx[key] = {"chat": c["id"],
                                "title": c.get("title", ""),
                                "q": um[:200].replace("\n", " "),
                                "a": am[:400].replace("\n", " "), "v": v}
                    added += 1
        if added:
            emb_save()
    finally:
        _emb_run.release()


def semantic_recall(query: str, exclude_id: str,
                    limit: int = 2) -> list[str] | None:
    """None means the embedding model is not available — the caller
    falls back to keyword matching rather than losing recall."""
    if not ensure_embedder(wait=3):
        return None
    index_exchanges()
    try:
        qv = embedder.embed(["search_query: " + query[:1000]])[0]
    except Exception:
        return None
    scored = sorted(((_cos(qv, e["v"]), e) for e in emb_entries()
                     if e.get("chat") != exclude_id and e.get("v")),
                    key=lambda t: -t[0])
    out = []
    for sim, e in scored[:limit]:
        if sim < EMBED_SIM_MIN:
            break
        out.append('From the conversation "%s" — asked: "%s" — answered: '
                   '"%s"' % (e["title"], e["q"], e["a"]))
    return out


@app.post("/api/embed/install")
def api_embed_install():
    if embed_model_file():
        threading.Thread(target=index_exchanges, daemon=True).start()
        return jsonify({"ok": True, "already": True})

    def run(task):
        files = engine.hf_files(cfg, EMBED_REPO)
        cand = [f for f in files if not f["split"] and not f["companion"]]
        pick = next((f for f in cand if f["quant"] == "Q8_0"), None) \
            or (min(cand, key=lambda f: f["size"]) if cand else None)
        if not pick:
            raise RuntimeError(f"No GGUF found in {EMBED_REPO}.")
        engine.download_sync(cfg, EMBED_REPO, pick["path"], EMBED_DIR,
                             task, 0, 90, label="recall model")
        if task.cancel:
            return
        task.set(pct=92, detail="Starting the embedding server…")
        if not ensure_embedder(30):
            raise RuntimeError("The embedding server did not come up.")
        task.set(pct=95, detail="Indexing past conversations…")
        index_exchanges()
        task.set(pct=100,
                 detail=f"Ready — {len(emb_index())} exchanges indexed.")

    return jsonify({"ok": True,
                    "task": engine.spawn("download", "recall model",
                                         run).view()})


# ---- Office documents: xlsx / docx / pptx built from tagged blocks ---- #
# The model writes a fenced block tagged xlsx, docx or pptx in a simple
# format (told to it via DOC_HOWTO); these builders turn that text into
# the real file. Libraries are imported lazily so the app still runs if
# they were not installed.
DOC_HOWTO = (
    "\n\nWhen the user asks for a spreadsheet, a Word document or a "
    "PowerPoint deck, put the content in ONE fenced code block tagged "
    "xlsx, docx or pptx. xlsx: lines of '# Sheet: <name>' each followed "
    "by CSV rows, first row the headers; after a sheet's rows you may add "
    "'# Chart: bar <title>' (or line, or pie) to embed a chart of that "
    "sheet — first column is the labels, the other columns the numbers. "
    "Chart types: bar to compare categories, line for change over time, "
    "stacked for parts within categories, scatter for two numeric "
    "columns, pie only for a few shares of a whole. Cells starting with "
    "= are real Excel formulas (e.g. =SUM(B2:B4), =B2*C2). docx: "
    "Markdown-style # headings, paragraphs, - bullets, **bold**, and "
    "Markdown tables (| cell | rows), which become real Word tables. "
    "pptx: optionally start with '# Template: <file>' to build on a "
    "company template from data/templates (its fonts, colours, layouts "
    "and master apply), or '# Theme: dark' / '# Theme: #RRGGBB'; then "
    "'# <slide title>' followed by - bullet lines per slide; a line "
    "'- image: <file in data/images>' (or 'image: latest' for the newest "
    "generated picture) places that image on the slide. The app offers "
    "the real file for download from that block.")


def deck_templates() -> list[str]:
    d = DATA_DIR / "templates"
    if not d.is_dir():
        return []
    return sorted(p.name for p in d.iterdir()
                  if p.suffix.lower() in (".pptx", ".potx"))


# Words that mean a conversation is about an Office file. Only the question
# itself is read — attached source files and search results ride along under
# their own headers and would match on any stray "slide" or "deck" in code.
DOC_ASK = re.compile(
    r"\b(xlsx|excel|spreadsheets?|workbooks?|docx|word\s+(?:doc|document|"
    r"file)s?|pptx|powerpoint|slides?|slide\s*decks?|decks?|"
    r"presentations?)\b", re.I)
DOC_FENCE = re.compile(r"```(?:xlsx|docx|pptx)\b", re.I)


def wants_documents(history: list[dict]) -> bool:
    """Whether the Office-file convention belongs in this request. It is
    ~400 tokens about CSV sheets and slide templates; sent with every coding
    question it was pure noise competing with the actual task."""
    for m in history[-6:]:
        text = m.get("content") or ""
        if m.get("role") == "user":
            ask = re.split(r"\n\n(?:Project files:|Web search results:)",
                           text, maxsplit=1)[0]
            if DOC_ASK.search(ask):
                return True
        elif DOC_FENCE.search(text):
            return True      # a follow-up to a file it already wrote
    return False


def doc_howto(history: list[dict] | None = None) -> str:
    if history is not None and not wants_documents(history):
        return ""
    names = deck_templates()
    if not names:
        return DOC_HOWTO
    return DOC_HOWTO + (" Available deck templates: " +
                        ", ".join(names[:6]) + ".")


def _load_template(fname: str) -> io.BytesIO:
    """A template from data/templates, as bytes python-pptx will open.
    .potx is the same zip as .pptx apart from its declared content type,
    so that one line is patched on the way in."""
    fname = (fname or "").strip()
    if "/" in fname or "\\" in fname or ".." in fname:
        raise RuntimeError("That template name is not allowed.")
    p = DATA_DIR / "templates" / fname
    if not p.is_file() or p.suffix.lower() not in (".pptx", ".potx"):
        avail = ", ".join(deck_templates()) or ("none yet — put .potx or "
                                                ".pptx files in "
                                                "data/templates")
        raise RuntimeError(f"No template called {fname}. Available: "
                           f"{avail}.")
    raw = p.read_bytes()
    if p.suffix.lower() == ".potx":
        import zipfile
        zin = zipfile.ZipFile(io.BytesIO(raw))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"presentationml.template.main+xml",
                        b"presentationml.presentation.main+xml")
                zo.writestr(item, data)
        raw = out.getvalue()
    return io.BytesIO(raw)


def _clear_slides(prs) -> None:
    """A .pptx used as a template may carry example slides; the deck
    starts clean while master, layouts and theme stay."""
    for sld in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sld.rId)
        prs.slides._sldIdLst.remove(sld)


def _pick_layout(prs):
    """The deck's 'title and content' layout: a title and a body (or
    content) placeholder. Any title plus any second placeholder used to
    do, and the Title Slide — title and subtitle — comes first in nearly
    every deck, so every slide came out as a title slide, bullets centred
    in the subtitle."""
    from pptx.enum.shapes import PP_PLACEHOLDER as PH
    titles = (PH.TITLE, PH.CENTER_TITLE)
    bodies = (PH.BODY, PH.OBJECT)
    layouts = list(prs.slide_layouts)
    kinds = [[ph.placeholder_format.type for ph in ly.placeholders]
             for ly in layouts]
    for ly, k in zip(layouts, kinds):
        if any(t in k for t in titles) and any(b in k for b in bodies):
            return ly
    for ly, k in zip(layouts, kinds):
        if any(t in k for t in titles) and \
                any(x not in titles and x != PH.SUBTITLE for x in k):
            return ly
    return layouts[1] if len(layouts) > 1 else layouts[0]


def build_xlsx(text: str) -> tuple[bytes, str]:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    # each sheet: (name, rows, chart) where chart is (type, title) or None
    sheets, name, rows, chart = [], "Sheet1", [], None
    for line in text.splitlines():
        m = re.match(r"\s*#+\s*Chart\s*:\s*(bar|line|pie|scatter|stacked)"
                     r"\b\s*(.*)$", line, re.I)
        if m:
            chart = (m.group(1).lower(), m.group(2).strip())
            continue
        m = re.match(r"\s*#+\s*Sheet\s*:\s*(.+)$", line, re.I)
        if m:
            if rows:
                sheets.append((name, rows, chart))
            # Excel refuses : / \ ? * [ ] in a sheet name, and quotes at
            # its ends; "Revenue: 2024" failed the whole export
            name = re.sub(r"[\\/*?:\[\]]", "-", m.group(1).strip())
            name, rows, chart = name.strip("'")[:31], [], None
        elif line.strip():
            rows.append(next(csv.reader([line])))
    if rows:
        sheets.append((name, rows, chart))
    if not sheets:
        raise RuntimeError("No rows found in the block.")
    for sname, srows, schart in sheets:
        ws = wb.create_sheet(title=sname or "Sheet")
        for r, row in enumerate(srows, 1):
            for c, val in enumerate(row, 1):
                v = val.strip()
                # a leading zero is an identifier (zip code, part no.),
                # not a number: "02134" stays text, never 2134
                if v and re.fullmatch(r"-?\d+(\.\d+)?", v) and \
                        not re.fullmatch(r"-?0\d+(\.\d+)?", v):
                    v = float(v) if "." in v else int(v)
                cell = ws.cell(row=r, column=c, value=v)
                if r == 1:
                    cell.font = Font(bold=True)
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value
                         is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = \
                min(40, width + 2)
        # An embedded chart of this sheet's data: labels from the first
        # column, values from the rest, series named by the header row.
        # Colors stay Excel's own theme — a fixed-order palette, never
        # hand-picked here.
        ncols = max((len(r) for r in srows), default=0)
        if schart and len(srows) >= 2 and ncols >= 2:
            from openpyxl.chart import ScatterChart, Series
            ctype, ctitle = schart
            nrows = len(srows)
            if ctype == "scatter":
                # x from the first column, one series per further column
                ch = ScatterChart()
                ch.style = 13
                xref = Reference(ws, min_col=1, min_row=2, max_row=nrows)
                for c in range(2, ncols + 1):
                    yref = Reference(ws, min_col=c, min_row=1,
                                     max_row=nrows)
                    ch.series.append(Series(yref, xref,
                                            title_from_data=True))
                ch.x_axis.title = str(srows[0][0]) if srows[0] else None
            else:
                if ctype == "stacked":
                    ch = BarChart()
                    ch.grouping = "stacked"
                    ch.overlap = 100
                else:
                    ch = {"bar": BarChart, "line": LineChart,
                          "pie": PieChart}[ctype]()
                last_col = 2 if ctype == "pie" else ncols
                data = Reference(ws, min_col=2, max_col=last_col,
                                 min_row=1, max_row=nrows)
                ch.add_data(data, titles_from_data=True)
                ch.set_categories(Reference(ws, min_col=1, min_row=2,
                                            max_row=nrows))
            ch.title = ctitle or sname
            ch.height, ch.width = 8, 15
            if ctype != "pie":
                # openpyxl 3.1 leaves the axes' delete flag unset, and
                # Excel then hides both axes: no labels, no scale
                ch.x_axis.delete = False
                ch.y_axis.delete = False
            ws.add_chart(ch, f"{get_column_letter(ncols + 2)}2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), ("application/vnd.openxmlformats-officedocument"
                            ".spreadsheetml.sheet")


def _docx_runs(paragraph, text: str) -> None:
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = (i % 2 == 1)


def build_docx(text: str) -> tuple[bytes, str]:
    from docx import Document
    doc = Document()
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue
        # A run of |-delimited lines is a Markdown table: header row bold,
        # the |---| separator dropped, rendered as a real Word table.
        if s.startswith("|") and s.endswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in
                         lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    tbl.append(cells)
                i += 1
            if tbl:
                cols = max(len(r) for r in tbl)
                t = doc.add_table(rows=len(tbl), cols=cols)
                t.style = "Table Grid"
                for r, rowvals in enumerate(tbl):
                    for c in range(cols):
                        para = t.cell(r, c).paragraphs[0]
                        txt = rowvals[c] if c < len(rowvals) else ""
                        if r == 0:
                            para.add_run(
                                re.sub(r"\*\*", "", txt)).bold = True
                        else:
                            _docx_runs(para, txt)
            continue
        m = re.match(r"(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(re.sub(r"\*\*", "", m.group(2)),
                            level=min(len(m.group(1)), 4))
        elif re.match(r"\s*[-*]\s+", line):
            _docx_runs(doc.add_paragraph(style="List Bullet"),
                       re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"\s*\d+[.)]\s+", line):
            _docx_runs(doc.add_paragraph(style="List Number"),
                       re.sub(r"^\s*\d+[.)]\s+", "", line))
        else:
            _docx_runs(doc.add_paragraph(), s)
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), ("application/vnd.openxmlformats-officedocument"
                            ".wordprocessingml.document")


def _pptx_image(name: str) -> Path | None:
    """Only images from the app's own gallery (data/images); 'latest'
    means the newest generated picture."""
    name = name.strip()
    if "/" in name or "\\" in name or ".." in name:
        return None
    if not IMAGES_OUT.is_dir():
        return None
    if name.lower() == "latest":
        pngs = sorted(IMAGES_OUT.glob("*.png"),
                      key=lambda p: -p.stat().st_mtime)
        return pngs[0] if pngs else None
    p = IMAGES_OUT / name
    return p if p.is_file() and p.suffix.lower() in (".png", ".jpg",
                                                     ".jpeg") else None


def build_pptx(text: str) -> tuple[bytes, str]:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    # optional first lines: '# Template: <file>' (wins over Theme) or
    # '# Theme: dark' / '# Theme: #RRGGBB'
    template = None
    bg = fg = accent = None
    slides, cur = [], None
    for raw in text.splitlines():
        s = raw.strip()
        if not s:
            continue
        m = re.match(r"#+\s*Template\s*:\s*(.+)$", s, re.I)
        if m:
            template = m.group(1).strip()
            continue
        m = re.match(r"#+\s*Theme\s*:\s*(\S+)", s, re.I)
        if m:
            t = m.group(1).lower()
            if t == "dark":
                bg, fg = RGBColor(0x1E, 0x1E, 0x1C), \
                    RGBColor(0xEC, 0xEA, 0xE4)
                accent = RGBColor(0xD9, 0x77, 0x57)
            elif re.fullmatch(r"#?[0-9a-f]{6}", t):
                h = t.lstrip("#")
                accent = RGBColor(int(h[0:2], 16), int(h[2:4], 16),
                                  int(h[4:6], 16))
            continue
        m = re.match(r"#+\s+(.*)", s)
        if m:
            cur = {"title": re.sub(r"\*\*", "", m.group(1)),
                   "bullets": [], "images": []}
            slides.append(cur)
            continue
        if cur is None:
            cur = {"title": "", "bullets": [], "images": []}
            slides.append(cur)
        body = re.sub(r"^[-*]\s+", "", s)
        m = re.match(r"(?:image|img)\s*:\s*(.+)$", body, re.I)
        if m:
            img = _pptx_image(m.group(1))
            if img:
                cur["images"].append(img)
            continue
        cur["bullets"].append(re.sub(r"\*\*", "", body))
    if not slides:
        raise RuntimeError("No slides found in the block.")
    if template:
        prs = Presentation(_load_template(template))
        _clear_slides(prs)
        bg = fg = accent = None      # the template owns the look
    else:
        prs = Presentation()
    layout = _pick_layout(prs)
    sw, sh = prs.slide_width, prs.slide_height
    for sl in slides:
        slide = prs.slides.add_slide(layout)
        if bg is not None:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = bg
        if slide.shapes.title is not None:
            slide.shapes.title.text = sl["title"][:90]
            for para in slide.shapes.title.text_frame.paragraphs:
                for run in para.runs:
                    if accent is not None:
                        run.font.color.rgb = accent
                    elif fg is not None:
                        run.font.color.rgb = fg
        body_ph = next((ph for ph in slide.placeholders
                        if ph.placeholder_format.idx != 0
                        and ph.has_text_frame), None)
        if body_ph is None:
            body_ph = slide.shapes.add_textbox(
                int(sw * 0.06), int(sh * 0.25),
                int(sw * 0.55), int(sh * 0.6))
        frame = body_ph.text_frame
        if sl["images"]:
            # Bullets keep the left, pictures take the right, whatever the
            # template's slide size. The box keeps the position and height
            # it inherits: setting the width alone wrote a box with no
            # position and no height.
            left, top, height = body_ph.left, body_ph.top, body_ph.height
            body_ph.left, body_ph.top = left, top
            body_ph.width, body_ph.height = int(sw * 0.48), height
        for i, btxt in enumerate(sl["bullets"][:12]):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = btxt[:200]
            for run in para.runs:
                if fg is not None:
                    run.font.color.rgb = fg
        imgs = sl["images"][:2]
        gap, top0 = int(sh * 0.03), int(sh * 0.2)
        box_w = int(sw * 0.41)
        box_h = (int(sh * 0.72) - gap * (len(imgs) - 1)) // max(1, len(imgs))
        for j, img in enumerate(imgs):
            # each picture scaled to fit its share of the right side, so a
            # square one never runs off the bottom of the slide
            pic = slide.shapes.add_picture(str(img), int(sw * 0.54),
                                           top0 + j * (box_h + gap),
                                           width=box_w)
            if pic.height > box_h:
                pic.width = int(pic.width * box_h / pic.height)
                pic.height = box_h
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), ("application/vnd.openxmlformats-officedocument"
                            ".presentationml.presentation")


DOC_BUILDERS = {"xlsx": build_xlsx, "docx": build_docx, "pptx": build_pptx}


@app.post("/api/doc")
def api_doc():
    b = request.get_json(silent=True)
    b = b if isinstance(b, dict) else {}
    kind = str(b.get("kind") or "").lower()
    if kind not in DOC_BUILDERS:
        return jsonify({"error": "Unknown document kind."}), 400
    title = re.sub(r"[^A-Za-z0-9 _-]+", "",
                   str(b.get("title") or "document")).strip()[:48] \
        or "document"
    try:
        data, mime = DOC_BUILDERS[kind](b.get("content") or "")
    except ImportError as exc:
        return jsonify({"error": f"The .{kind} exporter needs a Python "
                        f"package that is not installed ({exc}). Run: "
                        "pip install -r requirements.txt"}), 500
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": str(exc)}), 400
    return send_file(io.BytesIO(data), mimetype=mime, as_attachment=True,
                     download_name=f"{title}.{kind}")


# ---- images: local text-to-image via stable-diffusion.cpp ---- #
IMAGE_MODEL_REPO = "second-state/stable-diffusion-v1-5-GGUF"
IMAGE_MODEL_FILE = "stable-diffusion-v1-5-pruned-emaonly-Q8_0.gguf"
IMAGE_DIR = "image"
IMAGES_OUT = DATA_DIR / "images"


def image_model_file() -> Path | None:
    d = engine.MODELS_DIR / IMAGE_DIR
    if d.is_dir():
        for f in sorted(d.glob("*.gguf")):
            return f
    return None


@app.post("/api/image/install")
def api_image_install():
    if engine.sd_binary() and image_model_file():
        return jsonify({"ok": True, "already": True})

    def run(task):
        if not engine.sd_binary():
            engine.sd_install_sync(task, fit.hardware({}))
        if not image_model_file():
            engine.download_sync(cfg, IMAGE_MODEL_REPO, IMAGE_MODEL_FILE,
                                 IMAGE_DIR, task, 30, 100,
                                 label="image model")

    return jsonify({"ok": True,
                    "task": engine.spawn("download", "image engine",
                                         run).view()})


@app.post("/api/image/generate")
def api_image_generate():
    b = request.get_json(silent=True) or {}
    prompt = (b.get("prompt") or "").strip()
    if not prompt:
        return jsonify({"error": "Describe the image first."}), 400
    binary, model = engine.sd_binary(), image_model_file()
    if not binary or not model:
        return jsonify({"error": "The image engine is not installed yet — "
                                 "download it on the Images page."}), 400
    w = max(256, min(768, int(b.get("width") or 512))) // 64 * 64
    h = max(256, min(768, int(b.get("height") or 512))) // 64 * 64
    steps = max(1, min(50, int(b.get("steps") or 20)))
    seed = int(b.get("seed") or time.time()) % 2_000_000_000

    def run(task):
        IMAGES_OUT.mkdir(parents=True, exist_ok=True)
        name = f"img-{int(time.time())}-{seed}.png"
        out = IMAGES_OUT / name
        cmd = [str(binary), "-m", str(model), "-p", prompt,
               "-o", str(out), "--steps", str(steps),
               "-W", str(w), "-H", str(h), "-s", str(seed)]
        task.log(" ".join(cmd))
        proc = subprocess.Popen(cmd, stdout=subprocess.PIPE,
                                stderr=subprocess.STDOUT, text=True,
                                bufsize=1, cwd=str(engine.SD_DIR))
        for line in proc.stdout:
            line = line.strip()
            if not line:
                continue
            task.log(line)
            task.set(detail=line[:120])
            m = re.search(r"(\d+)\s*/\s*(\d+)", line)
            if m and int(m.group(2)) == steps:
                task.set(pct=int(m.group(1)) / steps * 100)
            if task.cancel:
                proc.kill()
                task.set(state="cancelled", detail="Cancelled")
                return
        proc.wait()
        if proc.returncode != 0 or not out.exists():
            raise RuntimeError(f"sd exited with code {proc.returncode} — "
                               "see the task log.")
        (IMAGES_OUT / (name + ".json")).write_text(
            json.dumps({"prompt": prompt, "seed": seed, "w": w, "h": h,
                        "steps": steps}), encoding="utf-8")
        task.meta["file"] = name
        task.set(pct=100, detail=f"Saved {name}")

    return jsonify({"ok": True,
                    "task": engine.spawn("image", prompt[:48],
                                         run).view()})


@app.get("/api/images")
def api_images_list():
    out = []
    if IMAGES_OUT.is_dir():
        for f in sorted(IMAGES_OUT.glob("*.png"),
                        key=lambda p: -p.stat().st_mtime)[:60]:
            meta = {}
            side = IMAGES_OUT / (f.name + ".json")
            if side.exists():
                try:
                    meta = json.loads(side.read_text(encoding="utf-8"))
                except Exception:
                    pass
            out.append({"name": f.name, "url": f"/images/{f.name}",
                        "prompt": meta.get("prompt", ""),
                        "created": f.stat().st_mtime})
    return jsonify({"images": out})


@app.get("/images/<name>")
def api_image_file(name: str):
    if "/" in name or "\\" in name or ".." in name:
        return jsonify({"error": "Bad name."}), 400
    return send_from_directory(IMAGES_OUT, name)


@app.delete("/api/images")
def api_image_delete():
    name = ((request.get_json(silent=True) or {}).get("name") or "")
    if "/" in name or "\\" in name or ".." in name or not name:
        return jsonify({"error": "Bad name."}), 400
    (IMAGES_OUT / name).unlink(missing_ok=True)
    (IMAGES_OUT / (name + ".json")).unlink(missing_ok=True)
    return jsonify({"ok": True})


# ---- voice: local speech-to-text via whisper.cpp ---- #
WHISPER_MODEL_REPO = "ggerganov/whisper.cpp"
WHISPER_MODEL_FILE = "ggml-base.bin"
VOICE_DIR = "voice"
whisper = engine.WhisperServer(8083)


def voice_model_file() -> Path | None:
    d = engine.MODELS_DIR / VOICE_DIR
    if d.is_dir():
        for f in sorted(d.glob("*.bin")):
            return f
    return None


def ensure_whisper(wait: int = 15) -> bool:
    model = voice_model_file()
    if not model or not engine.whisper_binary():
        return False
    if whisper.ready():
        return True
    if not whisper.alive():
        try:
            whisper.start_whisper(str(model))
        except Exception:
            return False
    return whisper.wait_ready(wait)


@app.post("/api/voice/install")
def api_voice_install():
    if engine.whisper_binary() and voice_model_file():
        ensure_whisper(20)
        return jsonify({"ok": True, "already": True})

    def run(task):
        if not engine.whisper_binary():
            engine.whisper_install_sync(task)
        if not voice_model_file():
            engine.download_sync(cfg, WHISPER_MODEL_REPO,
                                 WHISPER_MODEL_FILE, VOICE_DIR, task,
                                 45, 95, label="voice model")
            if task.cancel:
                return
        task.set(pct=96, detail="Starting the speech server…")
        if not ensure_whisper(30):
            raise RuntimeError("The whisper server did not come up.")
        task.set(pct=100, detail="Ready — speak with the mic button.")

    return jsonify({"ok": True,
                    "task": engine.spawn("download", "voice (whisper)",
                                         run).view()})


@app.post("/api/voice/transcribe")
def api_voice_transcribe():
    audio = request.get_data(cache=False)
    if not audio:
        return jsonify({"error": "No audio received."}), 400
    if len(audio) > 25_000_000:
        return jsonify({"error": "Recording too long."}), 400
    if not ensure_whisper(10):
        return jsonify({"error": "Voice input is not set up — download "
                                 "the Whisper model under Parameters."}), 503
    try:
        return jsonify({"text": whisper.transcribe(audio)})
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": str(exc)}), 502


_STOP = {"the", "and", "that", "this", "with", "from", "have", "what",
         "your", "about", "into", "does", "how", "can", "for", "you",
         "are", "was", "were", "will", "would", "could", "should", "did",
         "when", "where", "which", "there", "here", "then", "than",
         "please", "make", "just", "like", "also", "some", "them",
         "they", "their", "file", "files", "code"}


def _terms(text: str) -> set[str]:
    return {w for w in re.findall(r"[a-z0-9]{4,}", (text or "").lower())
            if w not in _STOP}


def recall_snippets(query: str, exclude_id: str, limit: int = 2) -> list[str]:
    """Excerpts from OTHER conversations that share enough distinct words
    with the question to plausibly matter. Keyword overlap, not magic —
    cheap, local, and inspectable."""
    q = _terms(query)
    if len(q) < 2:
        return []
    scored = []
    for c in read_chats():
        if c["id"] == exclude_id:
            continue
        msgs = c.get("messages", [])
        for i, m in enumerate(msgs):
            if m.get("role") != "user":
                continue
            reply = msgs[i + 1]["content"] if i + 1 < len(msgs) \
                and msgs[i + 1].get("role") == "assistant" else ""
            score = len(q & _terms(m.get("content", "") + " " + reply))
            if score >= 2:
                scored.append((score, c.get("title", ""),
                               m.get("content", ""), reply))
    scored.sort(key=lambda t: -t[0])
    out = []
    for _, title, um, am in scored[:limit]:
        out.append('From the conversation "%s" — asked: "%s" — answered: '
                   '"%s"' % (title, um[:200].replace("\n", " "),
                             am[:400].replace("\n", " ")))
    return out


def context_extra(query: str, chat_id: str) -> str:
    """What rides along with the system prompt: the remembered notes.

    Only what stays the same from turn to turn belongs here. llama.cpp
    keeps the prompt it has already read and reads again from the first
    token that differs; anything that changes with each question at the
    top made it re-read the whole conversation before every reply."""
    mem = read_memory().strip()
    parts = []
    if mem:
        parts.append("Things to remember from earlier (persistent memory; "
                     "the user can edit these):\n" +
                     newest_lines(mem, MEMORY_INJECT))
    if mem or cfg.get("recall", True):
        parts.append("To permanently remember a new important fact, put it "
                     "on its own line starting with remember: in your "
                     "reply.")
    return ("\n\n" + "\n\n".join(parts)) if parts else ""


def recall_extra(query: str, chat_id: str) -> str:
    """Excerpts from earlier conversations that bear on this question,
    when recall is on. They change with every question, so they go with
    the question itself, at the end of the prompt, not at its top."""
    if not cfg.get("recall", True):
        return ""
    snips = semantic_recall(query, chat_id)
    mode = "meaning"
    if snips is None:
        snips = recall_snippets(query, chat_id)
        mode = "keywords"
    if not snips:
        return ""
    # Background, not orders: a small model otherwise copies an old answer
    # that merely shares a few words with the new question.
    return ("Excerpts from earlier conversations (matched by " + mode +
            "). Use them only if they bear on the question below; "
            "otherwise ignore them:\n" + "\n".join(snips))


@app.get("/api/memory")
def api_memory_get():
    return jsonify({"text": read_memory(),
                    "recall": bool(cfg.get("recall", True))})


@app.post("/api/memory")
def api_memory_set():
    b = request.get_json(silent=True)
    if not isinstance(b, dict):
        return jsonify({"error": "Send the memory as {\"text\": …}."}), 400
    whole = write_memory(str(b.get("text") or ""))
    return jsonify({"ok": True, "trimmed": not whole,
                    "cap": MEMORY_CAP, "sent": MEMORY_INJECT})


# --------------------------------------------------------------------------- #
# engines (switchable model families in their own folders under ./models)
# --------------------------------------------------------------------------- #
ENGINE_ORDER = ["qwen", "coder", "gemma", "glm"]   # qwen first: the primary
SETUP_PLAN = [("qwen3-8b", "qwen"),       # (catalogue id, folder)
              ("gemma-4-e4b-unc", "gemma")]


def engines_list() -> list[dict]:
    """One entry per engine family folder that holds a finished model."""
    groups: dict[str, list[dict]] = {}
    for m in engine.local_models():
        # embedding and image models are helpers, never chat engines
        if m["folder"] and m["folder"] not in (EMBED_DIR, IMAGE_DIR) \
                and not m["partial"]:
            groups.setdefault(m["folder"], []).append(m)
    loaded_name = Path(server.model).name if server.model else ""
    out = []
    for fam in sorted(groups, key=lambda f: (ENGINE_ORDER.index(f)
                                             if f in ENGINE_ORDER else 99, f)):
        models = groups[fam]
        pick = next((m for m in models if m["name"] == loaded_name), None) \
            or next((m for m in models
                     if m["name"] == cfg.get("last_model")), None) \
            or models[0]
        out.append({"family": fam, "model": pick["name"],
                    "loaded": server.alive() and pick["name"] == loaded_name})
    return out


def model_runtime_args(kv_bits: int) -> list[str]:
    """llama-server flags corresponding to the memory profile we assess."""
    if int(kv_bits) == 8:
        return ["--cache-type-k", "q8_0", "--cache-type-v", "q8_0"]
    return []


# Only when nothing better can be read at all: flagged as estimated.
GUESSED_SHAPE = {"layers": 32, "kv_layers": 32, "kv_heads": 8,
                 "head_dim": 128, "exact": False, "hybrid": False,
                 "moe": False, "max_ctx": 0}


def model_shape(model: dict, online: bool = True) -> dict:
    """The architecture a file on disk is judged and placed by.

    config.json (the catalogue's, or the one saved at the last load when
    `online` is off) when it describes this very file — same layer count
    as the file's own header; otherwise the header itself, which is what
    llama.cpp will read. A catalogue name can match a bigger sibling, and
    a GGUF-only repo often has no config.json at all."""
    own = None if model.get("partial") else \
        fit.file_config(model.get("path") or "")
    conf = None
    if online:
        entry = fit.catalogue_match(model["name"])
        if entry:
            conf = fit.model_config(cfg, entry.get("config_repo", ""))
    else:
        conf = saved_model_config(model["name"])
    if conf and conf.get("exact") and (
            not own or int(conf["layers"]) == own["layers"]):
        return conf
    return own or conf or dict(GUESSED_SHAPE)


def saved_model_config(name: str) -> dict | None:
    """Return a usable persisted architecture, tolerating edited old config."""
    saved = cfg.get("last_model_config") or {}
    try:
        if saved.get("model") != name or not all(
                int(saved.get(key) or 0) > 0
                for key in ("layers", "kv_heads", "head_dim")):
            return None
    except (AttributeError, TypeError, ValueError):
        return None
    return {"layers": int(saved["layers"]),
            "kv_layers": int(saved.get("kv_layers") or saved["layers"]),
            "kv_heads": int(saved["kv_heads"]),
            "head_dim": int(saved["head_dim"]),
            "exact": bool(saved.get("exact")),
            "hybrid": bool(saved.get("hybrid")),
            "moe": bool(saved.get("moe")),
            "max_ctx": int(saved.get("max_ctx") or 0)}


def free_vram_hw(hw: dict) -> dict:
    """The card as Auto context should see it: only the VRAM that is free
    right now, measured, because the browser and the desktop are already
    using some of it. Call with our own server stopped, or it counts the
    model being replaced as taken."""
    free = int(hw.get("vram_free") or 0)
    if free and hw.get("vram") and not hw.get("planned"):
        return {**hw, "vram": min(int(hw["vram"]), free)}
    return hw


def resolve_ctx(size: int, conf: dict, hw: dict, bits: int,
                ctx: int | None = None, layout: dict | None = None) -> int:
    """The context a load uses: the one asked for, else the setting, and
    for Auto (0) whatever the fit calculation says costs nothing."""
    ctx = int(ctx or cfg.get("ctx") or 0)
    if ctx == 0:
        ctx = fit.auto_ctx(size, conf, free_vram_hw(hw), bits, layout)
    if not 512 <= ctx <= 262_144:
        raise RuntimeError("Context must be between 512 and 262144.")
    return ctx


def ngl_flag(gpu_layers: int, layers: int) -> int:
    """-ngl for llama-server. llama.cpp counts the output head as one layer
    past the last block, so `-ngl <layer count>` leaves that head — a
    vocabulary-sized matrix read every token — on the CPU. When the fit
    puts every layer on the card it has already counted the head's bytes
    (they are in the file size it divides up), so the head goes too."""
    return gpu_layers + 1 if gpu_layers >= layers else gpu_layers


def launch(model: dict, conf: dict, hw: dict, ctx: int, bits: int,
           gpu_layers: int | None = None, auto: bool = False) -> dict:
    """Start llama-server for a file on disk and wait until it answers.

    A mixture-of-experts model bigger than the card goes up with every
    layer on the GPU and the experts that do not fit held in system RAM
    (--n-cpu-moe, sized by fit.moe_plan from the file's tensor table). If
    that does not come up — a build older than the flag, or a compute
    buffer bigger than allowed for — the plain layer split the app always
    used is tried before giving up, so the speed-up can never cost a load.
    Anything else, and any deliberate gpu_layers, gets the plain split.

    `auto` says ctx came from Auto: a window picked for the expert
    placement is picked again by the layer rule if that placement fails."""
    def split(ctx: int) -> tuple[dict, int]:
        # Judged on the VRAM free now, as the window and the expert plan
        # are: on the whole card, a browser holding 1–2 GB made the split
        # ask for more than there was, and the load failed or spilled.
        a = fit.assess(model["size"], conf, free_vram_hw(hw), ctx, bits)
        requested = a["gpu_layers"] if gpu_layers is None \
            else int(gpu_layers)
        # A caller may deliberately offload fewer layers, but never exceed
        # the measured safe fit or the model's real layer count.
        return a, max(0, min(requested, a["gpu_layers"], a["layers"]))

    a, gl = split(ctx)
    plan = None
    if gpu_layers is None and gl < a["layers"]:
        plan = fit.moe_plan(fit.gguf_layout(model["path"]), conf,
                            free_vram_hw(hw), ctx, bits)
        if plan and not plan["n_cpu_moe"]:
            plan = None      # it all fits after all: nothing to hold back
    if plan:
        server.start(model["path"], ngl_flag(plan["layers"], plan["layers"]),
                     ctx, model_runtime_args(bits) +
                     ["--n-cpu-moe", str(plan["n_cpu_moe"])])
        if server.wait_ready(600):
            shown = {**a, **moe_fields(plan)}
            return {"fit": shown, "ctx": ctx, "gpu_layers": plan["layers"],
                    "layers": plan["layers"],
                    "n_cpu_moe": plan["n_cpu_moe"], "moe_fallback": False,
                    "why": fit.verdict_text(shown, hw)}
        if auto:
            ctx = fit.auto_ctx(model["size"], conf, free_vram_hw(hw), bits)
            a, gl = split(ctx)
    server.start(model["path"], ngl_flag(gl, a["layers"]), ctx,
                 model_runtime_args(bits))
    if not server.wait_ready(600):
        server.stop()
        raise RuntimeError("llama-server did not come up.")
    # the expert placement was tried and did not come up
    shown = {**a, **({"moe_layer_split": "fallback"} if plan else {})}
    return {"fit": shown, "ctx": ctx, "gpu_layers": gl,
            "layers": a["layers"], "n_cpu_moe": 0,
            "moe_fallback": bool(plan), "why": fit.verdict_text(shown, hw)}


def moe_fields(plan: dict) -> dict:
    """What a verdict needs to describe an expert placement. The file's own
    tensor table is the authority on MoE-ness, even when config.json could
    not be read — and a MoE gets no dense-model speed estimate."""
    return {"moe": True, "speed": None, "n_cpu_moe": plan["n_cpu_moe"],
            "expert_used": plan["expert_used"],
            "expert_count": plan["expert_count"]}


def loaded_sampling() -> dict:
    """The running model's own sampling profile, or {} when it publishes
    none (or none has been read yet)."""
    saved = cfg.get("last_model_sampling") or {}
    name = Path(server.model).name if server.model else ""
    if not name or saved.get("model") != name:
        return {}
    return {k: v for k, v in saved.items() if k != "model"}


def remember_sampling(name: str) -> dict:
    """Read a model's sampling from its repo and keep it beside its
    architecture, so an automatic start needs no network to use it. When the
    repo cannot be reached, what was saved for this model stands, and a model
    with nothing saved is left unrecorded so the next start asks again."""
    entry = fit.catalogue_match(name)
    repo = (entry or {}).get("config_repo") or (entry or {}).get("repo", "")
    profile = fit.model_sampling(cfg, repo) if repo else {}
    with config_lock:
        saved = cfg.get("last_model_sampling") or {}
        if profile is None:
            if saved.get("model") == name:
                return {k: v for k, v in saved.items() if k != "model"}
            cfg["last_model_sampling"] = {}
            save_config(cfg)
            return {}
        cfg["last_model_sampling"] = {"model": name, **profile}
        save_config(cfg)
    return profile


def sampling_params(b: dict) -> dict:
    """What the next reply samples with. A value in the request wins; then,
    unless the person switched to manual, the model's own settings; then
    the sliders. The Code page no longer forces a temperature of its own —
    0.2 is near-greedy, which thinking models are explicitly warned off."""
    out = {"temperature": cfg.get("temperature"), "top_p": cfg.get("top_p")}
    if cfg.get("sampling", "model") == "model":
        prof = loaded_sampling()
        for key in ("temperature", "top_p", "top_k", "min_p",
                    "repeat_penalty"):
            if prof.get(key) is not None:
                out[key] = prof[key]
    for key in ("temperature", "top_p"):
        if b.get(key) is not None:
            out[key] = b[key]
    return out


def load_model_by_name(name: str, ctx: int | None = None,
                       gpu_layers: int | None = None) -> dict:
    """Start llama-server using only measured hardware, never a plan card."""
    with engine_lock:
        model = next((m for m in engine.local_models()
                      if m["name"] == name), None)
        if not model:
            raise RuntimeError("That model is not on disk.")
        if model["partial"]:
            raise RuntimeError("That download has not finished.")
        if model.get("companion"):
            raise RuntimeError("That is a vision projector, a companion "
                               "file, not a model to load on its own.")
        chosen = int(ctx or 0)
        if chosen and not 512 <= chosen <= 262_144:
            raise RuntimeError("Context must be between 512 and 262144.")
        conf = model_shape(model)
        # The model being replaced goes before the card is measured, so the
        # free VRAM Auto sees is what the new one really has.
        server.stop()
        # plan_for is a catalogue comparison only. Using it here can select
        # too many layers, or even install flags for a GPU not in this PC.
        hw = fit.hardware({})
        bits = int(cfg.get("kv_bits", 8))
        auto = not chosen and not int(cfg.get("ctx") or 0)
        # A manual split gets the plain layer split, so its window is
        # picked by the layer rule too, not for an expert placement that
        # will not be used.
        ctx = resolve_ctx(model["size"], conf, hw, bits, chosen,
                          fit.gguf_layout(model["path"])
                          if gpu_layers is None else None)
        placed = launch(model, conf, hw, ctx, bits, gpu_layers, auto)
        ctx = placed["ctx"]
        with config_lock:
            cfg["last_model"] = name
            # Preserve the measured architecture so automatic startup uses
            # the same layer/KV calculation instead of a generic 32-layer
            # fallback (Qwen3 8B has 36 layers).
            cfg["last_model_config"] = {
                "model": name,
                **{key: conf.get(key) for key in
                   ("layers", "kv_layers", "kv_heads", "head_dim", "exact",
                    "hybrid", "moe", "max_ctx")},
            }
            # A size the caller named becomes the setting; Auto stays Auto
            # and only the window it resolved to is recorded.
            if chosen:
                cfg["ctx"] = ctx
            cfg["last_ctx"] = ctx
            save_config(cfg)
        remember_sampling(name)
        return {"model": name, **placed, "auto_ctx": auto,
                "sampling": loaded_sampling()}


def pick_quant(files: list[dict], conf: dict, hw: dict, ctx: int,
               kv_bits: int) -> dict | None:
    """The build first-run setup downloads: Q4_K_M when it runs here (the
    sane default quant), else the largest that fits VRAM, else the largest
    that at least fits RAM, else the smallest. Judged, never guessed."""
    cand = [f for f in files
            if not f["split"] and not f["companion"] and f["size"]]
    if not cand:
        return None
    scored = [(f, fit.assess(f["size"], conf, hw, ctx, kv_bits))
              for f in cand]
    q4 = next(((f, a) for f, a in scored if f["quant"] == "Q4_K_M"), None)
    if q4 and (q4[1]["verdict"] == "fits" or
               (not hw.get("vram") and q4[1]["fits_ram"])):
        return q4[0]
    fits = [f for f, a in scored if a["verdict"] == "fits"]
    if fits:
        return fits[-1]
    ram_ok = [f for f, a in scored if a["fits_ram"]]
    return ram_ok[-1] if ram_ok else cand[0]


@app.post("/api/engine/switch")
def api_engine_switch():
    fam = ((request.get_json(silent=True) or {}).get("family") or "").strip()
    eng = next((e for e in engines_list() if e["family"] == fam), None)
    if not eng:
        return jsonify({"error": f"No downloaded model for '{fam}' yet."}), 404
    if eng["loaded"]:
        return jsonify({"ok": True, "model": eng["model"], "already": True})
    try:
        out = load_model_by_name(eng["model"])
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": str(exc), "tail": server.tail(25)}), 500
    return jsonify({"ok": True, **out})


@app.post("/api/setup/firstrun")
def api_setup_firstrun():
    """Install llama.cpp if needed, download Qwen and Gemma into their own
    folders, then load Qwen as the primary engine."""
    running = next((t for t in engine.task_list()
                    if t.kind == "setup" and t.state == "running"), None)
    if running:
        return jsonify({"ok": True, "task": running.view()})

    def run(task):
        # First run judges against the machine actually here, never a
        # planned card.
        hw = fit.hardware({})
        if not engine.server_binary():
            task.set(detail="Installing llama.cpp…")
            engine.install_sync(task, hw)
        primary_file = ""
        spans = [(5, 50), (50, 95)]
        for (model_id, folder), (p0, p1) in zip(SETUP_PLAN, spans):
            if task.cancel:
                task.set(state="cancelled", detail="Cancelled — partial "
                         "downloads resume next time.")
                return
            entry = next(c for c in fit.CATALOGUE if c["id"] == model_id)
            task.set(pct=p0, detail=f"Choosing a {entry['name']} build…")
            files = engine.hf_files(cfg, entry["repo"])
            conf = fit.model_config(cfg, entry.get("config_repo")
                                    or entry["repo"])
            pick = pick_quant(files, conf, hw, int(cfg.get("ctx") or 8192),
                              int(cfg.get("kv_bits", 16)))
            if not pick:
                raise RuntimeError(f"No usable GGUF in {entry['repo']}.")
            task.log(f"{entry['name']}: {pick['name']} "
                     f"({pick['size']/1e9:.1f} GB) → models/{folder}/")
            dest = engine.download_sync(cfg, entry["repo"], pick["path"],
                                        folder, task, p0, p1,
                                        label=entry["name"])
            if task.cancel:
                return
            if model_id == SETUP_PLAN[0][0]:
                primary_file = dest.name
        # The tiny recall model rides along; skipping it never fails setup.
        if not embed_model_file():
            task.set(pct=95, detail="Fetching the small recall model…")
            try:
                files = engine.hf_files(cfg, EMBED_REPO)
                cand = [f for f in files
                        if not f["split"] and not f["companion"]]
                pick = next((f for f in cand if f["quant"] == "Q8_0"),
                            None) or min(cand, key=lambda f: f["size"])
                engine.download_sync(cfg, EMBED_REPO, pick["path"],
                                     EMBED_DIR, task, 95, 97,
                                     label="recall model")
            except Exception as exc:  # noqa: BLE001
                task.log(f"Recall model skipped: {exc}")
        task.set(pct=97, detail="Loading Qwen as the primary engine…")
        try:
            load_model_by_name(primary_file)
            task.log(f"Loaded {primary_file}.")
        except Exception as exc:  # noqa: BLE001
            task.log(f"Downloaded, but could not load the primary engine "
                     f"yet: {exc}")
        cfg["setup_complete"] = True
        save_config(cfg)
        task.set(pct=100, detail="Ready — Qwen is the primary engine; Gemma "
                                 "is one click away on the switcher.")

    task = engine.spawn("setup", "First-run setup", run)
    return jsonify({"ok": True, "task": task.view()})


# --------------------------------------------------------------------------- #
# fit
# --------------------------------------------------------------------------- #
@app.get("/api/fit")
def api_fit():
    """Every GGUF in a repo, each judged against this machine."""
    model_id = request.args.get("model", "")
    entry = next((m for m in fit.CATALOGUE if m["id"] == model_id), None)
    repo = request.args.get("repo") or (entry["repo"] if entry else "")
    if not repo:
        return jsonify({"error": "No model chosen."}), 400
    ctx = fit_ctx_arg()
    # Some entries exist for reference rather than for running: a serving
    # format this engine cannot load. Say why, and point at what does run.
    if entry and entry.get("runtime") and entry["runtime"] != "llama.cpp":
        hw = fit.hardware(cfg)
        need = entry.get("vram_needed") or 0
        alt = entry.get("gguf") or ""
        why = (f"{entry['name']} is a {entry['runtime']} format. llama.cpp "
               "loads GGUF only, so this build cannot run here.")
        if need:
            why += (f" It also wants about {need/1024**3:.0f} GB of VRAM "
                    "resident with no CPU offload, against your "
                    f"{hw.get('vram', 0)/1024**3:.0f} GB.")
        return jsonify({"repo": repo, "ctx": ctx, "hardware": hw,
                        "unsupported": True, "runtime": entry["runtime"],
                        "why": why, "alternative": alt, "files": [],
                        "recommended": ""})
    try:
        files = engine.hf_files(cfg, repo)
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": str(exc)}), 400
    hw = fit.hardware(cfg)
    conf = fit.model_config(cfg, (entry or {}).get("config_repo") or repo)
    out = []
    for f in files:
        a = assess_at(f["size"], conf, hw, ctx)
        out.append({**f, "fit": a, "why": fit.verdict_text(a, hw)})
    # Files come lowest quality first, so the best that fits is the last
    # one that does — not the first, which named Q2_K on almost any card.
    # Above Q8_0 is twice the size for no quality a reader would notice.
    fits = [f for f in out if f["fit"]["verdict"] == "fits"
            and not f["split"] and not f.get("companion")]
    sane = [f for f in fits if f["quant"] not in ("BF16", "F16", "F32")]
    best = (sane or fits or [None])[-1]
    return jsonify({"repo": repo, "ctx": ctx, "hardware": hw, "config": conf,
                    "files": out, "recommended": best["name"] if best else ""})


@app.get("/api/fit/local")
def api_fit_local():
    """The same judgement for a file already on disk."""
    hw = fit.hardware(cfg)
    ctx = fit_ctx_arg()
    out = []
    for m in engine.local_models():
        if m.get("folder") in (EMBED_DIR, IMAGE_DIR):
            continue
        if m["companion"]:
            out.append({**m, "fit": {}, "why": "vision projector — a "
                        "companion file for its model, not a model to "
                        "load on its own"})
            continue
        conf = model_shape(m)
        a = assess_at(m["size"], conf, hw, ctx,
                      "" if m["partial"] else m["path"])
        out.append({**m, "fit": a, "why": fit.verdict_text(a, hw)})
    return jsonify({"models": out, "hardware": hw, "ctx": ctx})


def fit_ctx_arg() -> int:
    """The context a fit page judges at: ?ctx= when given, else the
    setting. 0 is Auto — judged file by file in assess_at."""
    raw = request.args.get("ctx")
    try:
        return int(raw) if raw not in (None, "") else int(cfg.get("ctx") or 0)
    except ValueError:
        return int(cfg.get("ctx") or 0)


def assess_at(size: int, conf: dict, hw: dict, ctx: int,
              path: str = "") -> dict:
    """fit.assess at a fixed context, or on Auto at the one this file would
    get — so the fit column describes what loading it would really do. The
    whole card is the yardstick here, as for every other verdict on the
    page; the load itself also subtracts what other programs are using.

    A file already on disk also gets its exact expert placement, read from
    its own tensor table; one still on HuggingFace is described in general
    terms, since its header has not been downloaded."""
    bits = int(cfg.get("kv_bits", 16))
    layout = fit.gguf_layout(path) if path else None
    auto = ctx == 0
    if auto:
        ctx = fit.auto_ctx(size, conf, hw, bits, layout)
    a = {**fit.assess(size, conf, hw, ctx, bits), "auto_ctx": auto}
    if layout and layout["experts"] and a["gpu_layers"] < a["layers"]:
        plan = fit.moe_plan(layout, conf, hw, ctx, bits)
        if plan and plan["n_cpu_moe"]:
            a.update(moe_fields(plan))
        elif not plan:
            a.update({"moe": True, "speed": None, "moe_layer_split": "room"})
    return a


# --------------------------------------------------------------------------- #
# engine control
# --------------------------------------------------------------------------- #
def stop_engines() -> None:
    """Everything that runs from ./llama.cpp, stopped so a new build can
    take its place (Windows will not replace a DLL that is in use)."""
    with engine_lock:
        server.stop()
    embedder.stop()


@app.post("/api/engine/install")
def api_engine_install():
    try:
        return jsonify({"ok": True,
                        "task": engine.install(fit.hardware({}),
                                               stop_engines).view()})
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": str(exc)}), 400


@app.post("/api/engine/load")
def api_engine_load():
    b = request.get_json(silent=True) or {}
    name = (b.get("model") or "").strip()
    try:
        out = load_model_by_name(
            name, b.get("ctx"),
            b.get("gpu_layers") if b.get("gpu_layers") is not None else None)
    except Exception as exc:  # noqa: BLE001
        msg = str(exc)
        if "not on disk" in msg:
            return jsonify({"error": msg}), 404
        if "did not come up" in msg:
            return jsonify({"error": msg, "tail": server.tail(25)}), 500
        return jsonify({"error": msg}), 400
    return jsonify({"ok": True, **out})


@app.post("/api/engine/stop")
def api_engine_stop():
    server.stop()
    return jsonify({"ok": True})


@app.get("/api/engine/log")
def api_engine_log():
    return jsonify({"lines": server.tail(60), "running": server.alive(),
                    "ready": server.ready()})


# --------------------------------------------------------------------------- #
# models
# --------------------------------------------------------------------------- #
@app.get("/api/models")
def api_models():
    return jsonify({"local": engine.local_models(),
                    "dir": str(engine.MODELS_DIR)})


@app.post("/api/models/download")
def api_models_download():
    b = request.get_json(silent=True) or {}
    try:
        # Catalogue models with an engine family land in that family's own
        # folder, so the switcher finds them; anything else stays flat.
        entry = next((c for c in fit.CATALOGUE
                      if c["repo"] == b.get("repo")), None)
        folder = (entry or {}).get("family", "")
        task = engine.download(cfg, b.get("repo", ""), b.get("path", ""),
                               folder)
        return jsonify({"ok": True, "task": task.view()})
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": str(exc)}), 400


@app.delete("/api/models")
def api_models_delete():
    b = request.get_json(silent=True) or {}
    try:
        name = b.get("name", "")
        if server.model and Path(server.model).name == name:
            server.stop()
        engine.delete_model(name)
        return jsonify({"ok": True})
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": str(exc)}), 400


@app.get("/api/tasks")
def api_tasks():
    task_id = request.args.get("id", "")
    since = int(request.args.get("since", 0))
    if task_id:
        task = engine.TASKS.get(task_id)
        if not task:
            return jsonify({"error": "No such task."}), 404
        return jsonify(task.view(since))
    return jsonify([t.view(t.view()["cursor"]) for t in engine.task_list()[:20]])


@app.post("/api/tasks/<task_id>/cancel")
def api_task_cancel(task_id: str):
    task = engine.TASKS.get(task_id)
    if task:
        task.cancel = True
    return jsonify({"ok": True})


# --------------------------------------------------------------------------- #
# web search (opt-in, per message; DuckDuckGo's keyless HTML endpoint)
# --------------------------------------------------------------------------- #
SEARCH_URL = os.environ.get("LLAMA_STUDIO_SEARCH_URL",
                            "https://html.duckduckgo.com/html/")


def _ddg_target(href: str) -> str:
    """DuckDuckGo wraps result links in a redirect; unwrap to the real URL."""
    m = re.search(r"[?&]uddg=([^&]+)", href)
    return urllib.parse.unquote(m.group(1)) if m else href


@app.get("/api/search")
def api_search():
    q = (request.args.get("q") or "").strip()
    if not q:
        return jsonify({"error": "Nothing to search."}), 400
    try:
        r = requests.get(SEARCH_URL, params={"q": q}, timeout=12,
                         headers={"User-Agent": "Mozilla/5.0 (LlamaStudio)"})
        r.raise_for_status()
        page = r.text
        links = re.findall(
            r'class="result__a"[^>]*href="([^"]+)"[^>]*>(.*?)</a>',
            page, re.S)
        snippets = re.findall(r'class="result__snippet"[^>]*>(.*?)</a>',
                              page, re.S)

        def clean(s: str) -> str:
            return html_mod.unescape(re.sub(r"<[^>]+>", "", s)).strip()

        results = []
        for i, (href, title) in enumerate(links[:5]):
            results.append({"title": clean(title),
                            "url": _ddg_target(html_mod.unescape(href)),
                            "snippet": clean(snippets[i])
                            if i < len(snippets) else ""})
        if not results:
            return jsonify({"error": "The search engine returned nothing "
                                     "readable."}), 502
        return jsonify({"query": q, "results": results})
    except requests.RequestException as exc:
        return jsonify({"error": f"Search failed: {exc}"}), 502


# --------------------------------------------------------------------------- #
# workspace (a folder on this computer the Code page can read and edit)
# --------------------------------------------------------------------------- #
WS_SKIP = {".git", "node_modules", "__pycache__", ".venv", "venv", "env",
           "dist", "build", "target", ".idea", ".vscode", ".next",
           "site-packages"}
WS_MAX_FILES = 400
WS_MAX_READ = 300_000


def workspace_root() -> Path | None:
    raw = (cfg.get("workspace") or "").strip()
    if not raw:
        return None
    p = Path(raw).expanduser()
    return p if p.is_dir() else None


def ws_resolve(rel: str) -> Path:
    """A path inside the open folder — never outside it."""
    root = workspace_root()
    if not root:
        raise RuntimeError("No folder is open.")
    target = (root / rel).resolve()
    root = root.resolve()
    # by path parts, not a string prefix: a drive or filesystem root as
    # the folder doubled its separator ("//") and refused every file
    if target != root and root not in target.parents:
        raise RuntimeError("That path is outside the open folder.")
    return target


@app.post("/api/workspace")
def api_workspace_open():
    b = request.get_json(silent=True) or {}
    raw = (b.get("path") or "").strip()
    if not raw:
        cfg["workspace"] = ""
        save_config(cfg)
        return jsonify({"ok": True, "path": ""})
    p = Path(raw).expanduser()
    if not p.is_dir():
        return jsonify({"error": f"No such folder: {raw}"}), 400
    cfg["workspace"] = str(p)
    save_config(cfg)
    return jsonify({"ok": True, "path": str(p)})


@app.get("/api/workspace/tree")
def api_workspace_tree():
    root = workspace_root()
    if not root:
        return jsonify({"path": "", "files": []})
    files, truncated = [], False
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in sorted(dirnames)
                       if d not in WS_SKIP and not d.startswith(".")]
        rel_dir = Path(dirpath).relative_to(root)
        for name in sorted(filenames):
            if name.startswith(".") or name.endswith(".bak"):
                continue
            if len(files) >= WS_MAX_FILES:
                truncated = True
                break
            f = Path(dirpath) / name
            rel = str(rel_dir / name).replace("\\", "/").lstrip("./")
            try:
                files.append({"path": rel, "size": f.stat().st_size})
            except OSError:
                continue
        if truncated:
            break
    return jsonify({"path": str(root), "files": files,
                    "truncated": truncated})


@app.get("/api/workspace/file")
def api_workspace_read():
    try:
        target = ws_resolve(request.args.get("path", ""))
        if not target.is_file():
            raise RuntimeError("That file does not exist.")
        if target.stat().st_size > WS_MAX_READ:
            raise RuntimeError("That file is too large to attach "
                               f"({WS_MAX_READ // 1000} KB cap).")
        data = target.read_bytes()
        if b"\0" in data[:8192]:
            raise RuntimeError("That looks like a binary file.")
        return jsonify({"path": request.args.get("path", ""),
                        "content": data.decode("utf-8", errors="replace"),
                        "size": len(data)})
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": str(exc)}), 400


@app.post("/api/workspace/file")
def api_workspace_write():
    """Write model output back into the folder. The previous version is
    kept beside it as .bak — an edit is undoable, never silent loss."""
    b = request.get_json(silent=True) or {}
    try:
        target = ws_resolve(b.get("path") or "")
        content = b.get("content")
        if content is None:
            raise RuntimeError("Nothing to write.")
        backup = ""
        crlf = False
        if target.exists():
            old = target.read_bytes()
            crlf = b"\r\n" in old
            bak = target.with_name(target.name + ".bak")
            # a .bak that is a link would carry this write out of the folder
            bak.unlink(missing_ok=True)
            bak.write_bytes(old)
            backup = bak.name
        target.parent.mkdir(parents=True, exist_ok=True)
        # The file keeps the line endings it had. Text mode turned every
        # LF into CRLF on Windows, so one edit changed every line in git.
        text = str(content).replace("\r\n", "\n")
        if crlf:
            text = text.replace("\n", "\r\n")
        target.write_bytes(text.encode("utf-8"))
        return jsonify({"ok": True, "path": b.get("path"), "backup": backup})
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": str(exc)}), 400


@app.post("/api/workspace/run")
def api_workspace_run():
    """Run the user's own command inside the open folder and hand back
    exit code plus output. The command is always the one the person typed
    — model output never chooses what runs."""
    root = workspace_root()
    if not root:
        return jsonify({"error": "No folder is open."}), 400
    b = request.get_json(silent=True) or {}
    cmd = (b.get("command") or cfg.get("run_command") or "").strip()
    if not cmd:
        return jsonify({"error": "No run command set."}), 400
    cfg["run_command"] = cmd
    save_config(cfg)
    started = time.time()
    posix = os.name != "nt"
    # Its own process group, so what it started can be stopped with it (a
    # server the program launched kept running, and on Windows kept the
    # pipe open and the request hanging forever). Output goes to a file,
    # not a pipe: the command's own exit ends the wait, even when a child
    # it left behind still holds the output open. No stdin: a program
    # waiting on input() fails fast instead of waiting for a keyboard that
    # is not there. Read as UTF-8 whatever the console's code page.
    with tempfile.TemporaryFile() as log:
        try:
            p = subprocess.Popen(
                cmd, shell=True, cwd=str(root), stdin=subprocess.DEVNULL,
                stdout=log, stderr=subprocess.STDOUT,
                env={**os.environ, "PYTHONIOENCODING": "utf-8"},
                start_new_session=posix,
                creationflags=0 if posix else
                subprocess.CREATE_NEW_PROCESS_GROUP)
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": str(exc)}), 400
        try:
            p.wait(timeout=60)
            timed_out = False
        except subprocess.TimeoutExpired:
            timed_out = True
        kill_tree(p)          # all of it on a timeout, strays otherwise
        try:
            p.wait(timeout=10)
        except subprocess.TimeoutExpired:
            pass
        log.seek(0)
        out = log.read()[-64000:].decode("utf-8", "replace").strip()
    if timed_out:
        out = (out + "\n\n" if out else "") + "Timed out after 60 " \
            "seconds — stopped, with everything it started."
    return jsonify({"ok": True, "exit": -1 if timed_out else p.returncode,
                    "output": out[-8000:],
                    "seconds": round(time.time() - started, 1)})


def kill_tree(p: subprocess.Popen) -> None:
    """Stop a command and every process it started."""
    try:
        if os.name == "nt":
            subprocess.run(["taskkill", "/T", "/F", "/PID", str(p.pid)],
                           capture_output=True, timeout=15)
        else:
            os.killpg(p.pid, signal.SIGKILL)
    except Exception:  # noqa: BLE001
        try:
            p.kill()
        except Exception:  # noqa: BLE001
            pass


# --------------------------------------------------------------------------- #
# chat
# --------------------------------------------------------------------------- #
@app.get("/api/chats")
def api_chats():
    return jsonify([{k: v for k, v in c.items() if k != "messages"}
                    | {"count": len(c.get("messages", []))}
                    for c in read_chats()])


@app.get("/api/chats/<chat_id>")
def api_chat(chat_id: str):
    chat = get_chat(chat_id)
    return jsonify(chat) if chat else (jsonify({"error": "Not found."}), 404)


@app.delete("/api/chats/<chat_id>")
def api_chat_delete(chat_id: str):
    with chats_lock:
        write_chats([c for c in load_chats() if c["id"] != chat_id])
    # what it said stops being recalled into other conversations
    emb_forget(chat_id)
    return jsonify({"ok": True})


# Scaffolding the chat template wraps round the conversation, and the
# smallest reply worth leaving room for. Both are slack, not guesses at
# anything measurable: the prompt itself is counted by the tokenizer.
CTX_MARGIN = 160
PER_MSG = 6
REPLY_FLOOR = 512

CONTINUE_NUDGE = (
    "Your previous message was cut off before it finished. Carry straight "
    "on from the exact character it stopped at. Do not repeat any of it, do "
    "not start again, do not add a preamble or an apology — continue the "
    "text, and if it stopped inside a code block, continue the code and "
    "close the block properly."
)
# When the window ran out while the model was still thinking there is no
# answer to carry on from — only reasoning, which is never sent back.
ANSWER_NUDGE = (
    "You ran out of room while still thinking, before writing any answer. "
    "Write the answer now, directly and completely, without further "
    "deliberation."
)
# A continuation is mechanical — pick up at the last character — so it runs
# with thinking off. Left on, a thinking model re-deliberates the whole task
# from scratch in the little room left and patches in something different.
NO_THINK = {"enable_thinking": False}
# While a reply is being worked on the page hears from the server at least
# this often, text or no text, so a model that is reading a long prompt, or
# generating tokens that are not text, never looks like a frozen page. After
# QUIET seconds with no text, llama-server's /slots is asked what it is doing.
BEAT = 1.0
QUIET = 2.0


def turns(messages: list[dict]) -> list[dict]:
    """The stored conversation as the model is shown it: content only, with
    a question that never got an answer (Stop pressed before any text)
    folded into the next one, so the roles still alternate — several chat
    templates refuse two user turns in a row."""
    out: list[dict] = []
    for m in messages:
        if out and m["role"] == "user" and out[-1]["role"] == "user":
            out[-1] = {"role": "user", "content":
                       out[-1]["content"] + "\n\n" + m["content"]}
        else:
            out.append({"role": m["role"], "content": m["content"]})
    return out


def budget(system_text: str, history: list[dict], want: int,
           keep: int = 1) -> dict:
    """Fit prompt and reply inside the context window the server really has.

    Handed more prompt than it can hold, llama.cpp either refuses the request
    or slides its window along and drops the system prompt and the oldest
    turns without saying so — which from the outside looks exactly like the
    app forgetting the conversation mid-answer. So the trimming happens here
    instead: deliberate, counted with the model's own tokenizer, and
    reported back to the page.

    The last `keep` messages are never dropped: the question being asked,
    or for a continuation the question, the reply being carried on and the
    nudge. A continuation that lost its own reply carried on from nothing
    and appended unrelated text to the answer. If even those do not fit,
    the start of the reply is cut, since the model continues from its end.
    And what is sent always opens with the person: several chat templates
    (Gemma, Mistral) refuse a conversation that opens with the model.
    """
    n_ctx = server.n_ctx() or int(cfg.get("ctx") or 8192)
    want = max(0, int(want or 0))
    msgs = list(history)
    sys_t = server.count_tokens(system_text)
    fixed = sys_t + CTX_MARGIN + PER_MSG
    prompt = fixed + server.count_tokens(
        "\n".join(m["content"] for m in msgs)) + PER_MSG * len(msgs)
    dropped = 0
    if prompt > n_ctx - REPLY_FLOOR and len(msgs) > 1:
        # Only the path that needs per-message counts pays for them.
        counts = [server.count_tokens(m["content"]) + PER_MSG for m in msgs]
        keep = max(1, min(keep, len(msgs)))
        while len(msgs) > keep and fixed + sum(counts) > n_ctx - REPLY_FLOOR:
            msgs.pop(0)
            counts.pop(0)
            dropped += 1
        while len(msgs) > keep and msgs[0]["role"] != "user":
            msgs.pop(0)
            counts.pop(0)
            dropped += 1
        over = fixed + sum(counts) - (n_ctx - REPLY_FLOOR)
        long_one = max(range(len(msgs)), key=lambda i: counts[i])
        if over > 0 and msgs[long_one]["role"] == "assistant":
            text = msgs[long_one]["content"]
            per = len(text) / max(1, counts[long_one])
            cut = min(len(text), int((over + 32) * per * 1.1) + 1)
            msgs[long_one] = {**msgs[long_one], "content": "…" + text[cut:]}
            counts[long_one] = server.count_tokens(
                msgs[long_one]["content"]) + PER_MSG
        prompt = fixed + sum(counts)
    free = n_ctx - prompt
    cap = min(want, free) if want else free
    return {"history": msgs, "max_tokens": max(16, cap), "dropped": dropped,
            "n_ctx": n_ctx, "prompt_tokens": prompt, "free": free,
            "capped": bool(want and want < free), "overflow": free < 64}


def looks_unfinished(text: str, reason: str) -> bool:
    """Whether a reply stopped short of a whole answer. A token limit says so
    outright; an odd number of code fences says it just as plainly, since a
    block that was opened and never closed is a file that got cut in half."""
    if reason in ("length", "error"):
        return True
    body = text.rstrip()
    if not body:
        return False
    return body.count("```") % 2 == 1


@app.post("/api/chat")
def api_chat_send():
    """Streams the reply back as it is generated, then stores the exchange.

    With `continue: true` it resumes the last reply instead of starting a new
    one: the new text is appended to that same message, so a cut-off answer
    is finished in place rather than replaced by a second partial one.
    """
    b = request.get_json(silent=True) or {}
    text = (b.get("message") or "").strip()
    resume = bool(b.get("continue"))
    if not text and not resume:
        return jsonify({"error": "Nothing to send."}), 400
    if not server.ready():
        return jsonify({"error": "No model is loaded. Pick one on the Models "
                                 "page."}), 503

    chat_id = b.get("chat") or ""
    chat = get_chat(chat_id) if chat_id else None
    if resume:
        last = chat["messages"][-1] if chat and chat.get("messages") else {}
        written = (last.get("content") or "").strip()
        if last.get("role") != "assistant" or not (
                written or (last.get("reasoning") or "").strip()):
            return jsonify({"error": "There is no cut-off reply to "
                                     "continue."}), 400
        history = turns(chat["messages"])
        history.append({"role": "user",
                        "content": CONTINUE_NUDGE if written
                        else ANSWER_NUDGE})
        query = next((m["content"] for m in reversed(chat["messages"])
                      if m["role"] == "user"), "")
        # What is already written. Whether the answer is whole has to be
        # judged on the two halves joined — the closing fence arriving on
        # its own in the last round is a finished reply, not a new stump.
        prior = chat["messages"][-1]["content"]
    else:
        if not chat:
            chat = {"id": uuid.uuid4().hex[:12],
                    "title": " ".join(text.split()[:7])[:60] or "New chat",
                    "created": time.time(), "model": Path(server.model).name,
                    # which workspace this conversation belongs to: the plain
                    # chat page or the coding page. Old chats have no mode and
                    # are treated as plain chat.
                    "mode": "code" if b.get("mode") == "code" else "chat",
                    "messages": []}
        prior = ""
        chat["messages"].append({"role": "user", "content": text,
                                 "at": time.time()})
        # Stored before a single token comes back: if the person presses Stop
        # or closes the tab mid-reply, their own message is not lost with it.
        try:
            put_chat(chat)
        except OSError as exc:
            return jsonify({"error": "The conversation could not be saved "
                                     f"({exc}); nothing was sent."}), 503
        history = turns(chat["messages"])
        query = text

    system_text = ((b.get("system", cfg.get("system")) or "") +
                   context_extra(query, chat["id"]) +
                   doc_howto(history)).strip()
    if not resume:
        # sent with the question, never stored with it
        background = recall_extra(query, chat["id"])
        if background:
            history[-1] = {**history[-1], "content": background +
                           "\n\n---\n\n" + history[-1]["content"]}
    plan = budget(system_text, history,
                  b.get("max_tokens", cfg.get("max_tokens")),
                  keep=3 if resume else 1)
    params = {**sampling_params(b),
              "max_tokens": plan["max_tokens"], "system": system_text}
    if resume:
        params["chat_template_kwargs"] = NO_THINK

    def stream():
        started = time.time()
        pieces: list[str] = []
        thoughts: list[str] = []
        finished = False
        reason = ""
        failure = ""
        timings: dict = {}

        def rate(chars: int, seconds: float) -> float:
            # llama.cpp's own count of generated tokens per second when it
            # reports one; ~4 characters a token over the wall clock if not.
            measured = timings.get("predicted_per_second")
            if isinstance(measured, (int, float)) and measured > 0:
                return round(float(measured), 1)
            return round((chars / 4) / max(seconds, 0.001), 1)

        # llama-server is read on a thread of its own, so this generator can
        # speak every BEAT seconds even while the model sends nothing — and
        # by speaking, find out when the page has gone (Stop, a closed tab)
        # and hang up, which is what makes llama-server stop generating.
        handle: dict = {}
        events: queue.Queue = queue.Queue()

        def pump() -> None:
            try:
                for ev in server.chat_stream(plan["history"], params,
                                             handle=handle):
                    events.put(ev)
            except Exception as exc:  # noqa: BLE001
                if not handle.get("cancel"):
                    events.put({"error": str(exc)})
            finally:
                events.put(None)

        def alive(stage: str, prompt: dict, heard: float) -> dict:
            """What the model is doing, measured: llama.cpp's own prompt
            progress and token counts, never a guess at how long is left."""
            now = time.time()
            measured = handle.get("timings") or {}
            tokens = int(measured.get("predicted_n") or 0)
            tps = measured.get("predicted_per_second")
            if now - heard >= QUIET:
                slot = server.activity()
                tokens = max(tokens, int(slot.get("decoded") or 0))
                if stage == "wait" and slot.get("processing"):
                    stage = "prompt"
            if stage in ("wait", "prompt") and tokens:
                # generating, and not one of the tokens has been text
                stage = "hidden"
            out = {"stage": stage, "seconds": round(now - started),
                   "tokens": tokens, "quiet": round(now - heard)}
            if isinstance(tps, (int, float)) and tps > 0 and tokens:
                out["tps"] = round(float(tps), 1)
            if prompt.get("total"):
                out["prompt"] = [prompt["processed"], prompt["total"]]
            return {"alive": out}

        try:
            if not plan["overflow"]:
                threading.Thread(target=pump, daemon=True).start()
            yield "data: " + json.dumps(
                {"chat": chat["id"], "title": chat["title"],
                 "ctx": {"n_ctx": plan["n_ctx"],
                         "prompt": plan["prompt_tokens"],
                         "reply_room": plan["max_tokens"],
                         "dropped": plan["dropped"],
                         "capped": plan["capped"]}}) + "\n\n"
            if plan["overflow"]:
                yield "data: " + json.dumps(
                    {"error": "This conversation no longer fits in the "
                              "context window, even after trimming. Raise "
                              "the context size or start a new chat."}) \
                    + "\n\n"
            else:
                stage, prompt = "wait", {}
                heard = beat = started
                while True:
                    try:
                        ev = events.get(timeout=BEAT)
                    except queue.Empty:
                        ev = {}
                    if ev is None:
                        break
                    if "delta" in ev:
                        pieces.append(ev["delta"])
                        stage, heard = "write", time.time()
                        yield "data: " + json.dumps(
                            {"delta": ev["delta"]}) + "\n\n"
                    elif "think" in ev:
                        thoughts.append(ev["think"])
                        stage, heard = "think", time.time()
                        yield "data: " + json.dumps(
                            {"think": ev["think"]}) + "\n\n"
                    elif "prompt" in ev:
                        prompt = ev["prompt"]
                        if stage == "wait":
                            stage = "prompt"
                    elif "stop" in ev:
                        reason = ev["stop"]
                        timings.update(ev.get("timings") or {})
                    elif "error" in ev:
                        failure = ev["error"]
                        yield "data: " + json.dumps(
                            {"error": ev["error"]}) + "\n\n"
                    if time.time() - beat >= BEAT:
                        beat = time.time()
                        yield "data: " + json.dumps(
                            alive(stage, prompt, heard)) + "\n\n"
            finished = True
            if failure and not reason:
                # The stream died under the reply (llama-server crashed,
                # was unloaded, timed out): say so, as for any other end.
                reason = "error"
            elapsed = max(time.time() - started, 0.001)
            reply = "".join(pieces)
            thought = "".join(thoughts)
            produced = bool((prior + reply).strip() or thought.strip())
            yield "data: " + json.dumps(
                {"done": True,
                 "tps": rate(len(reply) + len(thought), elapsed),
                 "seconds": round(elapsed, 1),
                 "stop": reason,
                 "tokens": int(timings.get("predicted_n") or 0),
                 # all of the room went on thinking and none on the answer
                 "thought_only": bool(thought.strip() and not reply.strip()),
                 # no text at all came back, answer or reasoning
                 "empty": not produced,
                 # nothing written means nothing to carry on from
                 "unfinished": produced and looks_unfinished(
                     prior + reply, reason)}) + "\n\n"
        finally:
            # Page gone or reply over: hang up on llama-server either way,
            # so a Stop really stops the model instead of leaving it to
            # generate unseen up to the token limit.
            server.abort(handle)
            # Runs on a clean finish AND on GeneratorExit when the client goes
            # away, so a half-written reply is kept rather than thrown out.
            # One cut short still has llama.cpp's own speed to report.
            if not timings:
                timings.update(handle.get("timings") or {})
            reply = "".join(pieces)
            thought = "".join(thoughts)
            # Reasoning is kept for the reader but never sent back to the
            # model: history is built from content alone, as the thinking
            # models ask. A reply that is only reasoning is still stored, so
            # it can be carried on into an answer.
            if reply or thought.strip():
                elapsed = max(time.time() - started, 0.001)
                stop = reason if finished else "aborted"
                if resume:
                    # The continuation belongs to the message it continues,
                    # not beside it: one whole answer, not two halves.
                    last = chat["messages"][-1]
                    last["content"] = last["content"] + reply
                    if thought:
                        last["reasoning"] = (last.get("reasoning") or "") + \
                            thought
                    last["seconds"] = round(
                        float(last.get("seconds") or 0) + elapsed, 1)
                    last["tps"] = rate(len(last["content"]) + len(
                        last.get("reasoning") or ""),
                        float(last["seconds"]))
                    last["stop"] = stop
                    last["stopped"] = not finished
                    last["unfinished"] = looks_unfinished(
                        last["content"], stop)
                    last["continued"] = int(last.get("continued") or 0) + 1
                    if failure:
                        last["error"] = failure
                else:
                    msg = {"role": "assistant", "content": reply,
                           "at": time.time(),
                           "tps": rate(len(reply) + len(thought), elapsed),
                           "seconds": round(elapsed, 1),
                           "stopped": not finished, "stop": stop,
                           "unfinished": looks_unfinished(reply, stop)}
                    if thought:
                        msg["reasoning"] = thought
                    if failure:
                        msg["error"] = failure
                    chat["messages"].append(msg)
                chat["model"] = Path(server.model).name if server.model \
                    else chat.get("model")
                put_chat(chat)
                # Any `remember:` lines in the reply become permanent
                # memory for every future conversation, and the fresh
                # exchange joins the semantic index in the background.
                remember_lines(reply)
                threading.Thread(target=index_exchanges,
                                 daemon=True).start()

    return Response(stream(), mimetype="text/event-stream",
                    headers={"Cache-Control": "no-cache",
                             "X-Accel-Buffering": "no"})


# --------------------------------------------------------------------------- #
def autostart(model: dict) -> None:
    """Bring the last model back up at launch, placed exactly as a manual
    load would place it, from the architecture saved then — no network."""
    with engine_lock:
        # Planning mode must never affect what is launched on this PC.
        hw = fit.hardware({})
        conf = model_shape(model, online=False)
        bits = int(cfg.get("kv_bits", 8))
        try:
            ctx = resolve_ctx(model["size"], conf, hw, bits, None,
                              fit.gguf_layout(model["path"]))
            placed = launch(model, conf, hw, ctx, bits,
                            auto=not int(cfg.get("ctx") or 0))
            ctx = cfg["last_ctx"] = placed["ctx"]
            print(f"  {model['name']} ready at {ctx:,} context" +
                  (f", experts of {placed['n_cpu_moe']} blocks in RAM"
                   if placed["n_cpu_moe"] else "") + ".")
        except Exception as exc:  # noqa: BLE001
            print(f"  could not reload {model['name']}: {exc}")


def main() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    engine.MODELS_DIR.mkdir(parents=True, exist_ok=True)
    if cfg.get("auto_start") and cfg.get("last_model"):
        model = next((m for m in engine.local_models()
                      if m["name"] == cfg["last_model"] and not m["partial"]),
                     None)
        if model and engine.server_binary():
            # Waiting for the model (and, for a MoE, falling back to the
            # plain split if its placement does not start) happens beside
            # the web server, so the page opens straight away.
            print(f"  loading {model['name']}…")
            threading.Thread(target=autostart, args=(model,),
                             daemon=True).start()
            # Saved beside the architecture on every load; a model last
            # loaded by an older version has none yet, so read it once now,
            # off the startup path.
            if (cfg.get("last_model_sampling") or {}).get("model") != \
                    model["name"]:
                threading.Thread(target=remember_sampling,
                                 args=(model["name"],), daemon=True).start()
    url = f"http://127.0.0.1:{PORT}"
    print(f"\n  Llama Studio  →  {url}\n")
    if os.environ.get("LLAMA_STUDIO_NO_BROWSER") != "1":
        threading.Timer(1.2, lambda: webbrowser.open(url)).start()
    try:
        # This is a loopback-only desktop service, not an internet-facing web
        # deployment. Flask's threaded server also preserves streaming SSE.
        app.run(host="127.0.0.1", port=PORT, threaded=True, debug=False)
    finally:
        server.stop()
        embedder.stop()
        whisper.stop()


if __name__ == "__main__":
    main()
